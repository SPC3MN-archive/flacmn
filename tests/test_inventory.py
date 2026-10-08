import os
import re
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from iaflac import inventory as inv
from tests.helpers import HAVE_FLAC, encode_flac, tone_wav

THIN = Side(style="thin")
SHADE = PatternFill("solid", fgColor="FFF2F2F2")
HEADER_FILL = PatternFill("solid", fgColor="FF1F3864")
DATA = [   # already in the workbook's own order: artists case-SENSITIVELY sorted, albums in no particular order
    ("BROCKHAMPTON", "Saturation III", "Zipper", "MP3 (CBR 320kbps)", "MP3"),
    ("BROCKHAMPTON", "Saturation II", "Gold", "MP3 (CBR 320kbps)", "MP3"),
    ("Baby Keem", "The Melodic Blue", "trademark usa", "FLAC (16-bit/44.1kHz)", "FLAC"),
    ("Boards of Canada", "Music Has The Right to Children", "Wildlife Analysis", "FLAC (16-bit/44.1kHz)", "FLAC"),
    ("Boards of Canada", "Music Has The Right to Children", "An Eagle in Your Mind", "FLAC (16-bit/44.1kHz)", "FLAC"),
    ("MF Doom", "Mm.. Food", "Beef Rapp", "FLAC (16-bit/44.1kHz)", "FLAC"),
    ("Mac Demarco", "Salad Days", "Blue Boy", "MP3 (CBR 320kbps)", "MP3"),
    ("Mac Demarco", "Salad Days", "Brother", "FLAC (16-bit/44.1kHz)", "FLAC"),
    ("Radiohead", "OK Computer", "Airbag", "FLAC (16-bit/44.1kHz)", "FLAC"),
    ("TV Girl", "French Exit", "Birds Dont Sing", "FLAC (16-bit/44.1kHz)", "FLAC"),
    ("Talking Heads", "Remain in Light", "Once in a Lifetime", "FLAC (16-bit/44.1kHz)", "FLAC"),
    ("The Marias", "Cinema", "Hush", "FLAC (16-bit/44.1kHz)", "FLAC"),
    ("The Velvet Underground", "Loaded", "Sweet Jane", "FLAC (16-bit/44.1kHz)", "FLAC"),
]


def build_workbook(path: Path, rows=DATA, freeze="A2", broken_banding=False) -> Path:
    """A miniature of the real inventory: same sheets, headers, fonts, borders, banding and Summary layout."""
    wb = Workbook()
    lib = wb.active
    lib.title = "Library"
    summ = wb.create_sheet("Summary")

    def style(cell, parity, bold=False):
        cell.font = Font(name="Arial", size=10, bold=bold)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        if parity == 0:
            cell.fill = SHADE

    for ws, header, widths in ((lib, inv.LIBRARY_HEADER, (22, 34, 42, 26, 10)), (summ, inv.SUMMARY_HEADER, (24, 10, 10, 60))):
        for j, name in enumerate(header, 1):
            c = ws.cell(1, j, name)
            c.font, c.fill = Font(name="Arial", size=11, bold=True, color="FFFFFFFF"), HEADER_FILL
            c.alignment = Alignment(horizontal="center")
        ws.row_dimensions[1].height = 20
        for j, w in enumerate(widths):
            ws.column_dimensions["ABCDE"[j]].width = w
        ws.freeze_panes = freeze
        ws.sheet_view.zoomScale = 85
    for i, row in enumerate(rows):
        for j, v in enumerate(row, 1):
            style(lib.cell(i + 2, j, v), (i + 2) % 2)
    lib.auto_filter.ref = f"A1:E{1 + len(rows)}"
    computed = inv.summarise(list(rows))
    srows = sorted(((a, *v) for a, v in computed.items()), key=lambda r: r[0].casefold())
    for i, row in enumerate(srows):
        for j, v in enumerate(row, 1):
            style(summ.cell(i + 2, j, v), (i + 2) % 2)
    last = 1 + len(srows)
    summ.auto_filter.ref = f"A1:D{last}"
    for j, v in enumerate(("TOTAL", sum(r[1] for r in srows), sum(r[2] for r in srows)), 1):
        summ.cell(last + 2, j, v).font = Font(name="Arial", size=10, bold=True)
    if broken_banding:
        lib.cell(5, 1).fill = SHADE                     # row 5 is odd, so it should be plain
    wb.save(path)
    return path


def album(artist, album_name, songs, bits=16, rate=44100):
    rows = [(artist, album_name, s, inv.flac_file_type(bits, rate), "FLAC") for s in songs]
    return inv.NewAlbum(artist, album_name, rows)


class CleanSong(unittest.TestCase):
    def test_conventions_seen_in_the_real_inventory(self):
        cases = [
            ("1. Pink Floyd - Shine On You Crazy Diamond ( Part One ).flac", "Pink Floyd", "Shine On You Crazy Diamond ( Part One )"),
            ("01 - Speak To Me.flac", "Pink Floyd", "Speak To Me"),
            ("01 Airbag.flac", "Radiohead", "Airbag"),
            ("07. Come In Alone - My Bloody Valentine - Loveless (1991) (24bit FLAC).flac", "My Bloody Valentine",
             "Come In Alone - My Bloody Valentine - Loveless (1991) (24bit FLAC)"),
            ("02 - - Sherane a.k.a. Master Splinter's Daughter.flac", "Kendrick Lamar", "Sherane a.k.a. Master Splinter's Daughter"),
            ("05 Tell Him [live].flac", "Lauryn Hill", "Tell Him"),
            ("03 Street Spirit (Fade Out) [Acoustic].flac", "Radiohead", "Street Spirit (Fade Out)"),
            ("04 What Would Meek Do_.flac", "Pusha T", "What Would Meek Do"),
            ("Disc 1/01 Radiohead - Everything in Its Right Place.flac", "Radiohead", "Everything in Its Right Place"),
            ("01. Cygnus....Vismund Cygnus.flac", "The Mars Volta", "Cygnus....Vismund Cygnus"),
            ("01 - Miles Davis - So What.flac", "Miles Davis", "So What"),
            ("01 miles davis - So What.flac", "Miles Davis", "So What"),
            ("1-05 Disc Track.flac", "X", "Disc Track"),
        ]
        for filename, artist, expected in cases:
            self.assertEqual(inv.clean_song(filename, artist), expected, filename)

    def test_titles_that_are_numbers_survive_and_names_are_never_empty(self):
        self.assertEqual(inv.clean_song("745.flac", "Vince Staples"), "745")           # the real sheet has an empty cell here
        self.assertEqual(inv.clean_song("1979.flac", "Smashing Pumpkins"), "1979")
        self.assertEqual(inv.clean_song("03 1979.flac", "Smashing Pumpkins"), "1979")
        self.assertEqual(inv.clean_song("1000 Miles Away.flac", "Niontay"), "1000 Miles Away")
        self.assertEqual(inv.clean_song("[intro].flac", "X"), "[intro]")

    def test_file_type_strings(self):
        f = inv.flac_file_type
        self.assertEqual(f(16, 44100), "FLAC (16-bit/44.1kHz)")
        self.assertEqual(f(24, 96000), "FLAC (24-bit/96.0kHz)")
        self.assertEqual(f(24, 88200), "FLAC (24-bit/88.2kHz)")
        self.assertEqual(f(24, 48000), "FLAC (24-bit/48.0kHz)")
        self.assertEqual(f(16, 48000), "FLAC (16-bit/48.0kHz)")

    def test_summarise_matches_the_real_sheets_formats_column(self):
        rows = ([("A", "x", "s", "t", "MP3")] * 19 + [("A", "y", "s", "t", "FLAC")] * 7 + [("B", "z", "s", "t", "MP3")] * 5
                + [("B", "z", "s", "t", "FLAC")] * 5)
        out = inv.summarise(rows)
        self.assertEqual(out["A"], (2, 26, "MP3 (19), FLAC (7)"))
        self.assertEqual(out["B"], (1, 10, "FLAC (5), MP3 (5)"))       # ties: alphabetical


class WorkbookEditing(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.path = build_workbook(self.dir / "inv.xlsx")
        self.original_bytes = self.path.read_bytes()

    def tearDown(self):
        self.tmp.cleanup()

    def add(self, *albums, apply=True):
        book = inv.Inventory(self.path)
        plan = book.plan(list(albums))
        return book, plan, (book.apply(plan) if apply else None)

    def library(self):
        ws = load_workbook(self.path)["Library"]
        return [tuple(r) for r in ws.iter_rows(min_row=2, max_col=5, values_only=True)]

    def summary(self):
        ws = load_workbook(self.path)["Summary"]
        return [tuple(r) for r in ws.iter_rows(min_row=2, max_col=4, values_only=True)]

    # -- where things go ---------------------------------------------------------------------

    def test_a_new_artist_block_goes_where_the_case_sensitive_sort_puts_it(self):
        self.add(album("The Mars Volta", "Frances the Mute", ["Cygnus", "The Widow"]),
                 album("Miles Davis", "Kind of Blue", ["So What"]))
        artists = list(dict.fromkeys(r[0] for r in self.library()))
        self.assertEqual(artists, sorted(artists), "artist order must stay case-sensitively sorted")
        self.assertEqual(artists.index("Miles Davis"), artists.index("Mac Demarco") + 1)
        self.assertEqual(artists.index("The Mars Volta"), artists.index("The Marias") + 1)

    def test_summary_stays_case_insensitively_sorted_with_new_artists_in_place(self):
        self.add(album("The Mars Volta", "Frances the Mute", ["Cygnus"]), album("Miles Davis", "Kind of Blue", ["So What"]))
        names = [r[0] for r in self.summary() if r[0] not in (None, "TOTAL")]
        self.assertEqual(names, sorted(names, key=str.casefold))
        self.assertIn("Miles Davis", names)

    def test_existing_artist_reuses_its_spelling_and_the_album_joins_the_end_of_its_block(self):
        _, plan, _ = self.add(album("Boards Of Canada", "Geogaddi", ["Ready Lets Go", "Music Is Math"]))
        self.assertEqual(plan.placements[0].status, "existing artist")
        rows = self.library()
        block = [r for r in rows if r[0].casefold() == "boards of canada"]
        self.assertEqual({r[0] for r in block}, {"Boards of Canada"}, "no second spelling may appear")
        self.assertEqual([r[1] for r in block][-2:], ["Geogaddi", "Geogaddi"])
        self.assertEqual(len(block), 4)
        self.assertEqual(self.summary()[[r[0] for r in self.summary()].index("Boards of Canada")], ("Boards of Canada", 2, 4, "FLAC (4)"))

    def test_an_album_already_in_the_inventory_is_skipped_even_with_different_capitalisation(self):
        before = self.path.read_bytes()
        _, plan, result = self.add(album("boards of canada", "MUSIC HAS THE RIGHT TO CHILDREN", ["Wildlife Analysis"]))
        self.assertEqual(plan.placements[0].status, "already in inventory")
        self.assertFalse(result["written"])
        self.assertEqual(self.path.read_bytes(), before, "a no-op must not touch the file")
        self.assertEqual(list(self.dir.glob("*.backup-*")), [])

    def test_the_same_album_twice_is_added_once(self):
        a = album("Miles Davis", "Kind of Blue", ["So What"])
        book = inv.Inventory(self.path)
        plan = book.plan([a, a])
        self.assertEqual([p.status for p in plan.placements], ["new artist", "already in inventory"])

    # -- Summary and TOTAL -------------------------------------------------------------------

    def test_summary_rows_and_total_are_recomputed(self):
        _, plan, _ = self.add(album("Boards Of Canada", "Geogaddi", ["a", "b", "c"]), album("Miles Davis", "Kind of Blue", ["x", "y"]))
        rows = self.summary()
        by_artist = {r[0]: r for r in rows}
        self.assertEqual(by_artist["Boards of Canada"], ("Boards of Canada", 2, 5, "FLAC (5)"))
        self.assertEqual(by_artist["Miles Davis"], ("Miles Davis", 1, 2, "FLAC (2)"))
        self.assertEqual(by_artist["BROCKHAMPTON"], ("BROCKHAMPTON", 2, 2, "MP3 (2)"), "untouched artists are untouched")
        total = by_artist["TOTAL"]
        self.assertEqual(total[1:3], (sum(r[1] for r in rows if r[0] not in (None, "TOTAL")), len(self.library())))
        self.assertEqual(plan.totals_after, (total[1], total[2]))
        self.assertIsNone(rows[[r[0] for r in rows].index("TOTAL") - 1][0], "the blank row before TOTAL is kept")

    def test_a_consistent_workbook_reports_no_problems_and_a_tampered_one_does(self):
        self.assertEqual(inv.Inventory(self.path).consistency_problems(), [])
        wb = load_workbook(self.path)
        wb["Summary"]["B3"] = 99
        wb.save(self.path)
        self.assertTrue(inv.Inventory(self.path).consistency_problems())

    # -- look and feel -----------------------------------------------------------------------

    def test_banding_stays_correct_after_an_odd_number_of_inserted_rows(self):
        self.add(album("Miles Davis", "Kind of Blue", ["a", "b", "c"]))          # 3 rows: flips parity below the seam
        ws = load_workbook(self.path)["Library"]
        for r in range(2, ws.max_row + 1):
            for c in range(1, 6):
                self.assertEqual(ws.cell(r, c).fill.fill_type == "solid", r % 2 == 0, (r, c))
                self.assertEqual((ws.cell(r, c).font.name, ws.cell(r, c).font.sz, ws.cell(r, c).border.left.style), ("Arial", 10, "thin"))

    def test_summary_banding_font_and_total_style(self):
        self.add(album("Miles Davis", "Kind of Blue", ["a"]))
        ws = load_workbook(self.path)["Summary"]
        last = 1 + len([r for r in self.summary() if r[0] not in (None, "TOTAL")])
        for r in range(2, last + 1):
            self.assertEqual(ws.cell(r, 1).fill.fill_type == "solid", r % 2 == 0, r)
        self.assertTrue(ws.cell(last + 2, 1).font.b and ws.cell(last + 2, 1).value == "TOTAL")

    def test_autofilter_zoom_widths_freeze_and_header_are_preserved(self):
        self.add(album("Miles Davis", "Kind of Blue", ["a", "b"]))
        wb = load_workbook(self.path)
        lib, summ = wb["Library"], wb["Summary"]
        self.assertEqual(lib.auto_filter.ref, f"A1:E{lib.max_row}")
        n_artists = len([r for r in self.summary() if r[0] not in (None, "TOTAL")])
        self.assertEqual(summ.auto_filter.ref, f"A1:D{1 + n_artists}")
        self.assertEqual((lib.sheet_view.zoomScale, lib.freeze_panes, lib.column_dimensions["C"].width), (85, "A2", 42))
        self.assertEqual([c.value for c in lib[1]], inv.LIBRARY_HEADER)
        self.assertEqual(lib["A1"].font.color.rgb, "FFFFFFFF")
        self.assertEqual(lib.row_dimensions[1].height, 20)

    def test_every_original_row_is_still_there_in_order_and_unchanged(self):
        self.add(album("Miles Davis", "Kind of Blue", ["a"]), album("The Mars Volta", "Frances the Mute", ["b"]))
        it = iter(self.library())
        self.assertTrue(all(any(row == got for got in it) for row in DATA))

    # -- safety ------------------------------------------------------------------------------

    def test_a_backup_of_the_exact_original_is_made_and_no_temp_file_is_left(self):
        _, _, result = self.add(album("Miles Davis", "Kind of Blue", ["a"]))
        self.assertTrue(result["written"])
        self.assertEqual(result["backup"].read_bytes(), self.original_bytes)
        self.assertEqual([p.name for p in self.dir.iterdir() if "tmp" in p.name], [])

    def test_excel_lock_file_stops_the_write(self):
        (self.dir / "~$inv.xlsx").write_bytes(b"lock")
        book = inv.Inventory(self.path)
        with self.assertRaisesRegex(inv.InventoryError, "open"):
            book.apply(book.plan([album("Miles Davis", "Kind of Blue", ["a"])]))
        self.assertEqual(self.path.read_bytes(), self.original_bytes)

    def test_libreoffice_lock_file_also_stops_the_write(self):
        (self.dir / ".~lock.inv.xlsx#").write_bytes(b"lock")
        book = inv.Inventory(self.path)
        with self.assertRaisesRegex(inv.InventoryError, "open"):
            book.apply(book.plan([album("Miles Davis", "Kind of Blue", ["a"])]))
        self.assertEqual(self.path.read_bytes(), self.original_bytes)

    @unittest.skipIf(os.name == "nt", "POSIX permission bits")
    def test_file_permissions_survive_the_replace(self):
        self.path.chmod(0o640)
        self.add(album("Miles Davis", "Kind of Blue", ["a"]))
        self.assertEqual(self.path.stat().st_mode & 0o777, 0o640)

    def test_failed_verification_leaves_the_original_untouched(self):
        book = inv.Inventory(self.path)
        plan = book.plan([album("Miles Davis", "Kind of Blue", ["a"])])
        with mock.patch.object(inv, "verify_written", return_value=["simulated corruption"]):
            with self.assertRaisesRegex(inv.InventoryError, "nothing was changed"):
                book.apply(plan)
        self.assertEqual(self.path.read_bytes(), self.original_bytes)
        self.assertEqual([p.name for p in self.dir.iterdir()], ["inv.xlsx"], "no backup and no temp file may remain")

    def test_verification_really_catches_a_wrong_write(self):
        book = inv.Inventory(self.path)
        plan = book.plan([album("Miles Davis", "Kind of Blue", ["a"])])
        bad = self.dir / "bad.xlsx"
        shutil.copy(self.path, bad)                      # the original, i.e. "nothing was written"
        self.assertTrue(inv.verify_written(bad, plan, reference=self.path))

    def test_a_hand_edited_sheet_with_broken_banding_is_refused_not_guessed_at(self):
        broken = build_workbook(self.dir / "broken.xlsx", broken_banding=True)
        with self.assertRaisesRegex(inv.InventoryError, "banding|shading"):
            inv.Inventory(broken)

    def test_wrong_workbook_is_refused(self):
        wb = Workbook()
        wb.active.title = "Sheet1"
        p = self.dir / "other.xlsx"
        wb.save(p)
        with self.assertRaisesRegex(inv.InventoryError, "no 'Library' sheet"):
            inv.Inventory(p)


class StructuralValidator(unittest.TestCase):
    """ooxml_problems must be able to fail: each test breaks a good workbook in one specific way."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.good = build_workbook(self.dir / "good.xlsx")

    def tearDown(self):
        self.tmp.cleanup()

    def corrupt(self, part, transform=None, drop=False, name="bad.xlsx"):
        out = self.dir / name
        with zipfile.ZipFile(self.good) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.namelist():
                if item == part and drop:
                    continue
                data = zin.read(item)
                if item == part and transform:
                    data = transform(data.decode("utf8")).encode("utf8")
                zout.writestr(item, data)
        return out

    def test_a_good_workbook_is_clean(self):
        self.assertEqual(inv.ooxml_problems(self.good), [])

    def test_not_a_zip(self):
        bad = self.dir / "junk.xlsx"
        bad.write_bytes(b"this is not a workbook")
        self.assertTrue(any("zip" in p for p in inv.ooxml_problems(bad)))

    def test_malformed_xml(self):
        bad = self.corrupt("xl/worksheets/sheet1.xml", lambda x: x.replace("</worksheet>", ""))
        self.assertTrue(any("well-formed" in p for p in inv.ooxml_problems(bad)))

    def test_duplicate_cell(self):
        bad = self.corrupt("xl/worksheets/sheet1.xml", lambda x: x.replace('r="B2"', 'r="A2"', 1))
        self.assertTrue(any("out of order or duplicated" in p for p in inv.ooxml_problems(bad)))

    def test_rows_out_of_order(self):
        def swap(x):
            return x.replace('<row r="3"', '<row r="@"', 1).replace('<row r="4"', '<row r="3"', 1).replace('<row r="@"', '<row r="4"', 1)
        bad = self.corrupt("xl/worksheets/sheet1.xml", swap)
        self.assertTrue(inv.ooxml_problems(bad))

    def test_a_style_index_that_does_not_exist(self):
        bad = self.corrupt("xl/worksheets/sheet1.xml", lambda x: re.sub(r' s="\d+"', ' s="99"', x, count=1))
        self.assertTrue(any("uses style 99" in p for p in inv.ooxml_problems(bad)))

    def test_a_cell_in_the_wrong_row(self):
        bad = self.corrupt("xl/worksheets/sheet1.xml", lambda x: x.replace('r="A3"', 'r="A9"', 1))
        self.assertTrue(any("sits in row" in p for p in inv.ooxml_problems(bad)))

    def test_missing_part_referenced_by_a_relationship(self):
        bad = self.corrupt("xl/worksheets/sheet2.xml", drop=True)
        self.assertTrue(any("missing" in p for p in inv.ooxml_problems(bad)))

    def test_autofilter_that_disagrees_with_its_hidden_name(self):
        bad = self.corrupt("xl/worksheets/sheet1.xml", lambda x: re.sub(r'<autoFilter ref="[^"]*"', '<autoFilter ref="A1:E3"', x))
        self.assertTrue(any("autoFilter" in p for p in inv.ooxml_problems(bad)))

    def test_a_dimension_that_does_not_cover_the_data(self):
        bad = self.corrupt("xl/worksheets/sheet1.xml", lambda x: re.sub(r'<dimension ref="[^"]*"', '<dimension ref="A1:E3"', x))
        self.assertTrue(any("dimension" in p for p in inv.ooxml_problems(bad)))

    def test_apply_refuses_to_replace_the_file_when_the_written_package_is_structurally_broken(self):
        book = inv.Inventory(self.good)
        plan = book.plan([album("Miles Davis", "Kind of Blue", ["a"])])
        with mock.patch.object(inv, "ooxml_problems", return_value=["simulated structural damage"]):
            with self.assertRaisesRegex(inv.InventoryError, "nothing was changed"):
                book.apply(plan)
        self.assertEqual([p.name for p in self.dir.iterdir()], ["good.xlsx"])


class FromItemPreview(unittest.TestCase):
    def probe(self, bits, rate):
        from iaflac.flac_probe import FlacProbe, StreamInfo
        return FlacProbe(kind="flac", ok=True, streaminfo=StreamInfo(4096, 4096, 0, 0, rate, 2, bits, 1000, "ab" * 16))

    def test_preview_uses_names_and_headers_and_assumes_the_majority_for_unprobed_tracks(self):
        from iaflac.model import FileInfo, ItemInfo
        files = [FileInfo(f"Dir/{n:02d} - Miles Davis - Song {n}.flac") for n in (10, 2, 1, 3)]
        item = ItemInfo(identifier="x", files=files)
        probes = {files[1].name: self.probe(24, 96000), files[2].name: self.probe(24, 96000), files[3].name: self.probe(16, 44100)}
        rows = inv.album_from_item(item, probes, "Miles Davis", "Kind of Blue").rows
        self.assertEqual([r[2] for r in rows], ["Song 1", "Song 2", "Song 3", "Song 10"])
        self.assertEqual([r[3] for r in rows], ["FLAC (24-bit/96.0kHz)", "FLAC (24-bit/96.0kHz)", "FLAC (16-bit/44.1kHz)", "FLAC (24-bit/96.0kHz)"])

    def test_no_probed_headers_means_no_preview(self):
        from iaflac.model import ItemInfo
        with self.assertRaises(inv.InventoryError):
            inv.album_from_item(ItemInfo(identifier="x"), {}, "A", "B")


@unittest.skipUnless(HAVE_FLAC, "flac binary not installed")
class FromFolder(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def make(self, relpath, samples=4410, rate=44100):
        p = self.root / relpath
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(encode_flac(tone_wav(samples, rate=rate)))
        return p

    def test_rows_come_from_the_real_headers_in_track_order(self):
        for n in ("10 - Ten.flac", "02 - Two.flac", "01 - One.flac", "cover.jpg"):
            if n.endswith(".jpg"):
                (self.root / n).write_bytes(b"jpg")
            else:
                self.make(n)
        al = inv.album_from_folder(self.root, "Some Artist", "Some Album")
        self.assertEqual([r[2] for r in al.rows], ["One", "Two", "Ten"], "natural order, not 1, 10, 2")
        self.assertEqual({r[3] for r in al.rows}, {"FLAC (16-bit/44.1kHz)"})
        self.assertEqual({(r[0], r[1], r[4]) for r in al.rows}, {("Some Artist", "Some Album", "FLAC")})

    def test_multi_disc_albums_list_disc_1_before_disc_2(self):
        for n in ("Disc 2/01 Later.flac", "Disc 1/02 Second.flac", "Disc 1/01 First.flac", "Disc 10/01 Last.flac"):
            self.make(n)
        self.assertEqual([r[2] for r in inv.album_from_folder(self.root, "A", "B").rows], ["First", "Second", "Later", "Last"])

    def test_hires_file_type_is_read_from_the_file(self):
        self.make("01 Hi.flac", samples=9600, rate=96000)
        self.assertEqual(inv.album_from_folder(self.root, "A", "B").rows[0][3], "FLAC (16-bit/96.0kHz)")

    def test_a_fake_flac_is_refused(self):
        (self.root / "01 Fake.flac").write_bytes(b"ID3" + b"\0" * 200)
        with self.assertRaisesRegex(inv.InventoryError, "not a valid FLAC"):
            inv.album_from_folder(self.root, "A", "B")

    def test_non_flac_audio_and_empty_folders_are_refused(self):
        self.make("01 One.flac")
        (self.root / "02 Two.mp3").write_bytes(b"x")
        with self.assertRaisesRegex(inv.InventoryError, "non-FLAC"):
            inv.album_from_folder(self.root, "A", "B")
        empty = self.root / "empty"
        empty.mkdir()
        with self.assertRaisesRegex(inv.InventoryError, "no FLAC"):
            inv.album_from_folder(empty, "A", "B")


if __name__ == "__main__":
    unittest.main()
