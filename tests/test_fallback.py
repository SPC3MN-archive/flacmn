"""Vinyl rips and plain FLAC are searched alongside CD rips, and when no FLAC exists the search falls back to WAV and
VBR MP3: what is found, how it is tiered, and how it is saved (WAV encoded to FLAC, MP3 kept and labelled lossy)."""

import os
import tempfile
import unittest
from argparse import Namespace
from pathlib import Path
from unittest import mock

from mutagen.flac import FLAC
from openpyxl import load_workbook

from iaflac import altmedia, cli, download, preview, queries
from iaflac import inventory as inv
from iaflac.model import FileInfo, ItemInfo
from iaflac.pipeline import Hit, Options, Pipeline
from iaflac.queries import Want
from tests.helpers import HAVE_FFMPEG, HAVE_FLAC
from tests.test_cli_find import run_cli
from tests.test_inventory import DATA, build_workbook
from tests.test_pipeline import FakeIA, SAMPLES, make_alt_item, make_item

NEED = unittest.skipUnless(HAVE_FLAC and HAVE_FFMPEG, "flac and ffmpeg required")


def hit(identifier, rip):
    doc = {"identifier": identifier}
    return Hit(doc, "rung", None, rip)


class ChooseHits(unittest.TestCase):
    def test_cd_rips_cannot_fill_every_slot(self):
        plausible = [hit(f"cd{i}", True) for i in range(8)] + [hit("vinyl", False), hit("plain", False)]
        chosen = [h.doc["identifier"] for h in Pipeline.choose_hits(plausible, 6)]
        self.assertEqual(len(chosen), 6)
        self.assertIn("vinyl", chosen)
        self.assertIn("plain", chosen)
        self.assertEqual(chosen[:4], ["cd0", "cd1", "cd2", "cd3"], "the best CD rips keep their order")

    def test_nothing_changes_when_there_are_no_others_or_they_already_fit(self):
        only_cd = [hit(f"cd{i}", True) for i in range(8)]
        self.assertEqual(Pipeline.choose_hits(only_cd, 5), only_cd[:5])
        mixed = [hit("cd0", True), hit("plain", False), hit("cd1", True)]
        self.assertEqual(Pipeline.choose_hits(mixed, 3), mixed)


class SourceLabels(unittest.TestCase):
    def label(self, description="", files=(), title="Artist - Album", media="flac"):
        return preview.source_label(ItemInfo(identifier="x", title=title, description=description, media=media,
                                             files=[FileInfo(name=n, format=f, source="original") for n, f in files]))

    def test_a_rip_log_or_cue_means_a_cd_rip(self):
        self.assertEqual(self.label(files=[("a.flac", "Flac"), ("a.log", "Log")]), "CD rip (ships a rip log or cue)")

    def test_the_uploaders_own_words_are_reported_as_claims(self):
        self.assertEqual(self.label("Ripped from my LP on a turntable"), "vinyl rip (claimed by the uploader)")
        self.assertEqual(self.label("EAC secure mode rip"), "CD rip (claimed by the uploader)")
        self.assertEqual(self.label("bought on Bandcamp"), "web/store source (claimed by the uploader)")

    def test_no_source_stated_is_plain_flac(self):
        self.assertEqual(self.label("great album"), "plain FLAC (source not stated)")
        self.assertEqual(self.label("great album", media="wav"), "plain WAV (source not stated)")


@NEED
class VinylAndPlainAreIncluded(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.items = [make_item(f"cd-{i}", f"Test Artist - Test Album (CD {i})", "Test Artist", tracks=4, rip=True, downloads=500 - i, seed=i)
                     for i in range(1, 5)]
        cls.items += [
            make_item("lp-1", "Test Artist - Test Album [Vinyl Rip]", "Test Artist", tracks=4, rip=False, downloads=10, seed=9,
                      extra={"description": "24-bit needle drop from my turntable"}),
            make_item("plain-1", "Test Artist - Test Album (FLAC)", "Test Artist", tracks=4, rip=False, downloads=5, seed=10),
        ]

    def test_every_source_is_searched_alike_and_cd_rips_cannot_fill_every_slot(self):
        pipe = Pipeline(FakeIA(self.items))
        res = pipe.search_one(Want("Test Artist - Test Album", "Test Artist", "Test Album"), Options(max_candidates=4))
        labels = [label for label, _ in res.rungs]
        self.assertEqual(labels[0], "exact title")              # one query for any FLAC, whatever its source
        ids = [c.identifier for c in res.found]
        self.assertEqual(len(ids), 4)
        self.assertTrue({"lp-1", "plain-1"} & set(ids), "four CD rips would have filled every slot before")

    def test_the_listing_says_what_kind_of_upload_each_one_is(self):
        pipe = Pipeline(FakeIA(self.items))
        res = pipe.search_one(Want("Test Artist - Test Album", "Test Artist", "Test Album"), Options(max_candidates=6))
        text = preview.format_find_preview(res, 1, 1)
        self.assertIn("vinyl rip (claimed by the uploader)", text)
        self.assertIn("plain FLAC (source not stated)", text)
        self.assertIn("CD rip (ships a rip log or cue)", text)

    def test_rip_only_still_means_cd_rips_only(self):
        res = Pipeline(FakeIA(self.items)).search_one(Want("Test Artist - Test Album", "Test Artist", "Test Album"),
                                                      Options(rip_only=True, max_candidates=10))
        self.assertEqual({c.identifier for c in res.found}, {f"cd-{i}" for i in range(1, 5)})

    def test_vinyl_and_plain_flac_get_a_real_tier_and_can_be_tested_and_saved(self):
        pipe = Pipeline(FakeIA(self.items))
        opts = Options(max_candidates=6, deep="none", spectral=False)
        res = pipe.search_one(Want("Test Artist - Test Album", "Test Artist", "Test Album"), opts)
        pipe.verify_selected(res, [c for c in res.found if c.identifier in ("lp-1", "plain-1")], opts)
        tiers = {c.identifier: c.tier for c in res.candidates}
        self.assertTrue(tiers, "tested")
        self.assertTrue(all(t in "ABC" for t in tiers.values()), tiers)
        self.assertEqual(cli.savable([res])[1], [], "none of them is blocked")


@NEED
class Fallback(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wav = make_alt_item("wav-1", "Test Artist - Test Album", "Test Artist", "wav", seed=1)
        cls.vbr = make_alt_item("vbr-1", "Test Artist - Test Album", "Test Artist", "mp3", seed=2)
        cls.flac = make_item("flac-1", "Test Artist - Something Different", "Test Artist", tracks=4, seed=3, album_tag="Something Different")
        cls.want = Want("Test Artist - Test Album", "Test Artist", "Test Album")

    def search(self, items, **opts):
        opts = Options(**{"fallback": ("wav", "vbr"), **opts})
        return Pipeline(FakeIA(items)).search_one(self.want, opts)

    def test_no_flac_means_wav_and_vbr_mp3_are_searched_with_wav_first(self):
        res = self.search([self.vbr, self.wav, self.flac])
        self.assertEqual([(c.identifier, c.item.media) for c in res.found], [("wav-1", "wav"), ("vbr-1", "mp3")])
        self.assertTrue(any("WAV" in label for label, _ in res.rungs))
        self.assertTrue(any("no FLAC found" in n for n in res.notes))

    def test_the_wav_track_list_ignores_ias_derived_mp3s(self):
        res = self.search([self.wav])
        self.assertEqual([f.name for f in res.found[0].item.flacs], [f"0{i} Track {i}.wav" for i in range(1, 5)])

    def test_a_real_flac_match_means_no_fallback_at_all(self):
        res = self.search([make_item("flac-2", "Test Artist - Test Album", "Test Artist", tracks=4, seed=4), self.wav])
        self.assertEqual([c.identifier for c in res.found], ["flac-2"])
        self.assertEqual([c.item.media for c in res.found], ["flac"])

    def test_the_fallback_can_be_limited_or_turned_off(self):
        self.assertEqual([c.identifier for c in self.search([self.wav, self.vbr], fallback=("wav",)).found], ["wav-1"])
        vbr_only = self.search([self.wav, self.vbr], fallback=("vbr",)).found
        self.assertEqual({(c.identifier, c.item.media) for c in vbr_only}, {("wav-1", "mp3"), ("vbr-1", "mp3")},
                         "asked for VBR only: a WAV upload is listed through its derived MP3s")
        off = self.search([self.wav, self.vbr], fallback=())
        self.assertEqual(off.found, [])
        self.assertIn("no FLAC candidate matched this request", off.notes)

    def test_rip_only_never_falls_back(self):
        self.assertEqual(self.search([self.wav], rip_only=True).found, [])

    def test_the_listing_says_these_are_wav_or_lossy(self):
        text = preview.format_find_preview(self.search([self.wav, self.vbr]), 1, 1)
        self.assertIn("4 WAV (WAV, converted to FLAC when saved)", text)
        self.assertIn("4 MP3 (VBR MP3, lossy)", text)
        self.assertIn("plain WAV (source not stated)", text)

    def test_tiers_for_the_new_media(self):
        pipe = Pipeline(FakeIA([self.wav, self.vbr]))
        opts = Options(fallback=("wav", "vbr"), deep="none")
        res = pipe.search_one(self.want, opts)
        pipe.verify_selected(res, list(res.found), opts)
        tier = {c.identifier: c.tier for c in res.candidates}
        self.assertEqual(tier["wav-1"], "B", "CD format and every track a whole number of sectors")
        wav = next(c for c in res.candidates if c.identifier == "wav-1")
        self.assertEqual(wav.alt_stats.aligned, 4)
        self.assertIn("V1", wav.levels)
        self.assertIn("WAV", " ".join(wav.provenance.reasons))

    def test_a_wav_that_is_not_sector_aligned_is_only_tier_c_and_lossy_mp3_is_tier_l(self):
        odd = make_alt_item("wav-odd", "Test Artist - Test Album", "Test Artist", "wav", seed=5, samples=SAMPLES + 7)
        pipe = Pipeline(FakeIA([odd, self.vbr]))
        opts = Options(fallback=("wav", "vbr"), deep="none")
        res = pipe.search_one(self.want, opts)
        pipe.verify_selected(res, list(res.found), opts)
        primary = res.candidates[0]
        self.assertEqual(primary.tier, "C")
        alts = {a["identifier"]: a["tier"] for a in primary.alternates}
        self.assertEqual(alts.get("vbr-1"), "L")

    def test_files_that_are_not_what_they_claim_are_tier_d_and_never_offered(self):
        fake = make_alt_item("wav-fake", "Test Artist - Test Album", "Test Artist", "wav", seed=6)
        for name in list(fake.files):
            if name.endswith(".wav"):
                fake.files[name] = fake.files[name.replace(".wav", ".mp3")]      # an MP3 renamed .wav
        pipe = Pipeline(FakeIA([fake]))
        opts = Options(fallback=("wav",), deep="none")
        res = pipe.search_one(self.want, opts)
        pipe.verify_selected(res, list(res.found), opts)
        self.assertEqual(res.candidates[0].tier, "D")
        self.assertIn("not real WAV", " ".join(res.candidates[0].provenance.reasons))
        self.assertEqual(cli.savable([res])[0], [])


class HeaderProbes(unittest.TestCase):
    def test_a_compressed_wav_is_refused(self):
        import struct
        head = b"RIFF" + struct.pack("<I", 100) + b"WAVE" + b"fmt " + struct.pack("<IHHIIHH", 16, 2, 2, 44100, 0, 4, 4) + b"data" + struct.pack("<I", 4) + b"xxxx"
        p = altmedia.probe_wav(head, len(head))
        self.assertFalse(p.ok)
        self.assertIn("not PCM", p.detail)

    def test_not_a_wav_at_all_is_named_by_what_it_is(self):
        self.assertEqual(altmedia.probe_wav(b"fLaC" + bytes(100), 104).kind, "flac")
        self.assertEqual(altmedia.probe_wav(b"OggS" + bytes(100), 104).kind, "ogg")

    def test_a_streamed_header_with_no_length_uses_the_file_size(self):
        import struct
        head = (b"RIFF" + struct.pack("<I", 0) + b"WAVE" + b"fmt " + struct.pack("<IHHIIHH", 16, 1, 2, 44100, 176400, 4, 16)
                + b"data" + struct.pack("<I", 0xFFFFFFFF) + bytes(4 * 588 * 3))
        p = altmedia.probe_wav(head, len(head))
        self.assertTrue(p.ok)
        self.assertEqual(p.samples, 588 * 3)
        self.assertTrue(p.sector_aligned)


class Planning(unittest.TestCase):
    def item(self, media):
        files = [FileInfo("01 - Artist - One.wav", "WAVE", "original", 10, md5="0" * 32), FileInfo("02 - Artist - Two.wav", "WAVE", "original", 10, md5="1" * 32),
                 FileInfo("01 - Artist - One.mp3", "VBR MP3", "derivative", 5), FileInfo("cover.jpg", "JPEG", "original", 3)]
        return ItemInfo(identifier="x", title="Artist - Album", files=files, media=media)

    def rels(self, res):
        return sorted(str(pf.path.relative_to(res.folder.resolve())) for pf in res.planned)

    def test_a_wav_album_downloads_the_wavs_and_names_the_flacs_the_library_way(self):
        res = download.plan_library(self.item("wav"), Path("/tmp/iaflac-lib"), "Artist", "Album", rip="skip", track_names="library", media="wav")
        self.assertEqual(self.rels(res), ["01 One.wav", "02 Two.wav", "cover.jpg"])
        self.assertEqual([pf.final.name for pf in res.planned if pf.final], ["01 One.flac", "02 Two.flac"])

    def test_a_vbr_album_takes_the_vbr_files_even_though_ia_calls_them_derivatives(self):
        res = download.plan_library(self.item("mp3"), Path("/tmp/iaflac-lib"), "Artist", "Album", rip="skip", track_names="library", media="mp3")
        self.assertEqual(self.rels(res), ["01 One.mp3", "cover.jpg"])
        self.assertEqual([pf.final for pf in res.planned], [None, None])


@NEED
class SaveTheFallback(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wav = make_alt_item("wav-1", "Test Artist - Test Album", "Test Artist", "wav", seed=1)
        cls.vbr = make_alt_item("vbr-1", "Test Artist - Test Album", "Test Artist", "mp3", seed=2)

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        self.out, self.lib, self.trash = str(root / "runs"), root / "1Music", root / "trash"
        self.trash.mkdir()
        (self.lib / "Radiohead").mkdir(parents=True)
        self.xlsx = build_workbook(self.lib / "1Music_Library_Inventory.xlsx")
        self.original = self.xlsx.read_bytes()
        self.env = mock.patch.dict(os.environ, {"IAFLAC_CONFIG": str(root / "config.json"), "IAFLAC_TRASH": str(self.trash)})
        self.env.start()
        for k in ("IAFLAC_LIBRARY", "IAFLAC_INVENTORY"):
            os.environ.pop(k, None)

    def tearDown(self):
        self.env.stop()
        self.tmp.cleanup()

    def args(self, *items, **kw):
        base = dict(items=list(items or ["Test Artist - Test Album"]), file=None, top=3, include_live=False, free_only=False,
                    deep="none", no_spectral=True, max_deep_mb=1500, out=self.out, rip_only=False, max_candidates=10,
                    search_only=False, select=None, library=str(self.lib), inventory=None, no_inventory=False, save=None,
                    yes=False, art="cover", rip_files="skip", track_names="library", replace=False, scan=False, fallback="both")
        return Namespace(**{**base, **kw})

    def album(self):
        return self.lib / "Test Artist" / "Test Album"

    def rows(self):
        return [r for r in load_workbook(self.xlsx)["Library"].iter_rows(min_row=2, max_col=5, values_only=True) if r[0] == "Test Artist"]

    def test_a_wav_album_is_downloaded_encoded_to_flac_tagged_and_inventoried_as_flac(self):
        code, out, err, asked = run_cli(cli.cmd_find, self.args(), FakeIA([self.wav]), answers=["1", "y"], interactive=True)
        self.assertEqual(code, 0, err)
        self.assertIn("(WAV, will be converted to FLAC)", asked[1]["shown"])
        self.assertIn("converted to FLAC (verified as it is converted)", asked[1]["shown"])
        self.assertEqual(sorted(os.listdir(self.album())), [f"0{i} Track {i}.flac" for i in range(1, 5)], "no WAV, no .part left")
        tags = FLAC(self.album() / "02 Track 2.flac")
        self.assertEqual((tags["artist"], tags["album"], tags["title"], tags["tracknumber"]),
                         (["Test Artist"], ["Test Album"], ["Track 2"], ["2"]))
        self.assertEqual({r[3:5] for r in self.rows()}, {("FLAC (16-bit/44.1kHz)", "FLAC")})
        self.assertEqual(len(self.rows()), 4)
        self.assertEqual(inv.Inventory(self.xlsx).consistency_problems(), [])

    def test_the_flac_matches_the_wav_audio_exactly(self):
        import subprocess
        run_cli(cli.cmd_find, self.args(), FakeIA([self.wav]), answers=["1", "y"], interactive=True)
        decoded = subprocess.run(["flac", "-d", "-s", "-c", str(self.album() / "01 Track 1.flac")], capture_output=True).stdout
        self.assertEqual(decoded, self.wav.files["01 Track 1.wav"])

    def test_a_vbr_mp3_album_is_saved_as_mp3_labelled_lossy_and_inventoried_as_mp3(self):
        code, out, err, asked = run_cli(cli.cmd_find, self.args(), FakeIA([self.vbr]), answers=["1", "y"], interactive=True)
        self.assertEqual(code, 0, err)
        self.assertIn("** LOSSY: VBR MP3 **", asked[1]["shown"])
        self.assertEqual(sorted(os.listdir(self.album())), [f"0{i} Track {i}.mp3" for i in range(1, 5)])
        rows = self.rows()
        self.assertEqual({r[4] for r in rows}, {"MP3"})
        self.assertTrue(all(r[3].startswith("MP3 (VBR ") and r[3].endswith("kbps)") for r in rows), rows)
        self.assertEqual(inv.Inventory(self.xlsx).consistency_problems(), [])

    def test_fallback_none_finds_nothing_and_saves_nothing(self):
        code, out, _, asked = run_cli(cli.cmd_find, self.args(fallback="none"), FakeIA([self.wav, self.vbr]), answers=["1"], interactive=True)
        self.assertEqual(asked, [], "nothing to choose from")
        self.assertFalse(self.album().exists())

    def test_a_failed_encode_keeps_the_wav_and_leaves_no_half_album_in_the_inventory(self):
        real = download.subprocess.run

        def failing(cmd, *a, **k):
            if cmd and Path(str(cmd[0])).stem == "flac" and "--best" in cmd:
                class R:
                    returncode, stderr, stdout = 1, b"ERROR: boom", b""
                return R()
            return real(cmd, *a, **k)

        with mock.patch("iaflac.download.subprocess.run", failing):
            code, out, err, _ = run_cli(cli.cmd_find, self.args(), FakeIA([self.wav]), answers=["1", "y"], interactive=True)
        self.assertEqual(code, 1)
        self.assertEqual(self.rows(), [])
        self.assertEqual(self.xlsx.read_bytes(), self.original)
        self.assertTrue(list(self.album().glob("*.wav")), "the downloaded WAVs are kept")
        self.assertTrue((self.album() / ".iaflac-incomplete").exists(), "still marked incomplete, so a rerun resumes")

    def test_saving_a_wav_album_twice_does_nothing_the_second_time(self):
        run_cli(cli.cmd_find, self.args(select="1", save="1", yes=True), FakeIA([self.wav]), interactive=False)
        first = self.xlsx.read_bytes()
        client = FakeIA([self.wav])
        code, out, err, _ = run_cli(cli.cmd_find, self.args(select="1", save="1", yes=True), client, interactive=False)
        self.assertEqual(code, 0, err)
        self.assertEqual(self.xlsx.read_bytes(), first)
        self.assertEqual([c for c in client.calls if c[0] == "stream"], [], "no WAV was downloaded again")

    # ---- replace ---------------------------------------------------------------------------------

    def old_mp3_album(self):
        old = self.album()
        old.mkdir(parents=True)
        for i in range(1, 5):
            (old / f"0{i} Old {i}.mp3").write_bytes(b"mp3")
        return old

    def test_replace_uses_wav_and_turns_an_mp3_album_into_flac(self):
        old = self.old_mp3_album()
        args = self.args("Test Artist/Test Album", fallback="wav")
        code, out, err, asked = run_cli(cli.cmd_replace, args, FakeIA([self.wav, self.vbr]), answers=["1", "y"], interactive=True)
        self.assertEqual(code, 0, err)
        self.assertEqual(list(old.glob("*.mp3")), [])
        self.assertEqual(len(list(old.glob("*.flac"))), 4)
        self.assertEqual(len(list(self.trash.iterdir())), 1)

    def test_replace_never_offers_a_lossy_replacement_for_an_mp3_album(self):
        old = self.old_mp3_album()
        args = self.args("Test Artist/Test Album", fallback="wav")
        run_cli(cli.cmd_replace, args, FakeIA([self.vbr]), interactive=False)
        self.assertEqual(len(list(old.glob("*.mp3"))), 4)
        self.assertEqual(list(self.trash.iterdir()), [])

    def test_a_lossy_candidate_can_not_be_forced_into_a_replacement(self):
        old = self.old_mp3_album()
        client = FakeIA([self.vbr])
        cand = Pipeline(client).load("vbr-1", media="mp3")
        code = cli.library_flow(client, self.lib, None, [cli.Job("vbr-1", None, cand, old)], yes=False, art="cover",
                                rip="skip", track_names="library")
        self.assertEqual(code, 2)


if __name__ == "__main__":
    unittest.main()
