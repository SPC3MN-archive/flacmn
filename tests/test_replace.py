"""Replacing an album you already have (typically MP3) with a verified FLAC version: the inventory rows, the folder
swap and its safety properties, and the whole flow through the real CLI handlers."""

import os
import tempfile
import unittest
from argparse import Namespace
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from iaflac import cli, replace
from iaflac import inventory as inv
from tests.helpers import HAVE_FFMPEG, HAVE_FLAC
from tests.test_cli_find import run_cli
from tests.test_inventory import DATA, build_workbook
from tests.test_pipeline import FakeIA, make_item

MP3_ROWS = [("Test Artist", "Test Album", f"Old Song {i}", "MP3 (CBR 128kbps)", "MP3") for i in range(1, 6)]


def new_album(artist, album, n, replaces=""):
    return inv.NewAlbum(artist, album, [(artist, album, f"Track {i}", "FLAC (16-bit/44.1kHz)", "FLAC") for i in range(1, n + 1)],
                        replaces=replaces)


class InventoryReplacement(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        rows = sorted(DATA + MP3_ROWS, key=lambda r: r[0])
        self.path = build_workbook(Path(self.tmp.name) / "inv.xlsx", rows=rows)
        self.book = inv.Inventory(self.path)

    def tearDown(self):
        self.tmp.cleanup()

    def test_the_old_rows_are_swapped_for_the_new_ones_in_the_same_place(self):
        plan = self.book.plan([new_album("Test Artist", "Test Album", 4, replaces="Test Album")])
        self.assertEqual(plan.placements[0].status, "replaces existing rows")
        self.assertEqual(len(plan.removed_rows), 5)
        self.assertEqual(plan.rows_after, plan.rows_before - 5 + 4)
        self.assertEqual([r[3] for r in plan.library_rows if r[0] == "Test Artist"], ["FLAC (16-bit/44.1kHz)"] * 4)
        first_old = next(i for i, r in enumerate(self.book.library_rows) if r[0] == "Test Artist")
        first_new = next(i for i, r in enumerate(plan.library_rows) if r[0] == "Test Artist")
        self.assertEqual(first_old, first_new, "same position in the sheet")
        summary = {r[0]: r for r in plan.summary_rows}["Test Artist"]
        self.assertEqual(summary[1:], (1, 4, "FLAC (4)"))

    def test_the_written_workbook_verifies_and_only_the_replaced_rows_are_gone(self):
        plan = self.book.plan([new_album("Test Artist", "Test Album", 4, replaces="Test Album")])
        self.book.apply(plan)
        again = inv.Inventory(self.path)
        self.assertEqual(again.consistency_problems(), [])
        self.assertEqual([r for r in again.library_rows if r[0] == "Test Artist" and r[4] == "MP3"], [])
        self.assertEqual(len([r for r in again.library_rows if r[0] == "Radiohead"]), 1, "other artists untouched")
        self.assertEqual(len(list(Path(self.tmp.name).glob("*.backup-*.xlsx"))), 1)

    def test_a_replacement_may_use_a_different_album_name_than_the_old_rows(self):
        plan = self.book.plan([new_album("Test Artist", "Test Album (Deluxe)", 4, replaces="Test Album")])
        self.assertEqual({r[1] for r in plan.library_rows if r[0] == "Test Artist"}, {"Test Album (Deluxe)"})

    def test_replacing_with_identical_rows_is_a_no_op(self):
        self.book.apply(self.book.plan([new_album("Test Artist", "Test Album", 4, replaces="Test Album")]))
        again = inv.Inventory(self.path)
        plan = again.plan([new_album("Test Artist", "Test Album", 4, replaces="Test Album")])
        self.assertFalse(plan.changes)

    def test_if_the_old_album_is_not_in_the_inventory_the_new_one_is_simply_added(self):
        plan = self.book.plan([new_album("Test Artist", "Other Album", 3, replaces="No Such Album")])
        self.assertEqual(plan.placements[0].status, "existing artist")
        self.assertEqual(plan.removed_rows, [])
        self.assertEqual(plan.rows_after, plan.rows_before + 3)

    def test_rows_that_are_not_together_are_not_replaced_automatically(self):
        rows = sorted(DATA + MP3_ROWS, key=lambda r: r[0])
        i = next(k for k, r in enumerate(rows) if r[0] == "Test Artist")
        rows.insert(i + 2, ("Test Artist", "Different Album", "x", "FLAC (16-bit/44.1kHz)", "FLAC"))
        book = inv.Inventory(build_workbook(Path(self.tmp.name) / "split.xlsx", rows=rows))
        with self.assertRaises(inv.InventoryError):
            book.plan([new_album("Test Artist", "Test Album", 4, replaces="Test Album")])


class FolderSwap(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.lib = Path(self.tmp.name) / "lib"
        self.trash = Path(self.tmp.name) / "trash"
        self.trash.mkdir()
        self.old = self.lib / "Artist" / "Album"
        self.old.mkdir(parents=True)
        for i in (1, 2, 3):
            (self.old / f"0{i} Song.mp3").write_bytes(b"mp3")
        (self.old / "cover.jpg").write_bytes(b"old cover")
        (self.old / "booklet.pdf").write_bytes(b"pdf")
        (self.old / ".DS_Store").write_bytes(b"x")

    def tearDown(self):
        self.tmp.cleanup()

    def staging(self, marker=False, cover=True):
        st = replace.staging_folder(self.old, "item-1")
        st.mkdir()
        for i in (1, 2, 3):
            (st / f"0{i} Song.flac").write_bytes(b"flac")
        if cover:
            (st / "cover.jpg").write_bytes(b"new cover")
        if marker:
            (st / ".iaflac-incomplete").write_text("item-1\n")
        return st

    def test_survey_counts_audio_by_type_and_ignores_hidden_files(self):
        info = replace.survey(self.old)
        self.assertEqual(dict(info.audio), {".mp3": 3})
        self.assertTrue(info.replaceable)
        self.assertEqual(info.describe(), "3 MP3")

    def test_an_album_with_any_flac_is_not_replaceable(self):
        (self.old / "04 Bonus.flac").write_bytes(b"f")
        self.assertFalse(replace.survey(self.old).replaceable)

    def test_scan_lists_only_albums_with_audio_and_no_flac(self):
        done = self.lib / "Artist" / "Done"
        done.mkdir()
        (done / "01 a.flac").write_bytes(b"f")
        (self.lib / "Artist" / "Empty").mkdir()
        self.assertEqual([a.path.name for a in replace.scan(self.lib)], ["Album"])

    def test_find_existing_matches_close_names_and_the_album_dash_artist_style(self):
        (self.lib / "Artist" / "Other - Artist").mkdir()
        (self.lib / "Artist" / "Other - Artist" / "01 a.mp3").write_bytes(b"x")
        self.assertEqual(replace.find_existing(self.lib, "artist", "album").path, self.old)
        self.assertEqual(replace.find_existing(self.lib, "Artist", "Other").path.name, "Other - Artist")
        self.assertIsNone(replace.find_existing(self.lib, "Artist", "Something Else Entirely"))
        self.assertIsNone(replace.find_existing(self.lib, "Nobody", "Album"))

    def test_volume_numbers_must_agree(self):
        (self.lib / "Artist" / "Greatest Hits Vol. 1").mkdir()
        self.assertIsNone(replace.find_existing(self.lib, "Artist", "Greatest Hits Vol. 2"))

    def test_resolve_folder_accepts_artist_slash_album_and_refuses_anything_outside_the_library(self):
        self.assertEqual(replace.resolve_folder("Artist/Album", self.lib), self.lib / "Artist" / "Album")
        with self.assertRaises(replace.ReplaceError):
            replace.resolve_folder("Artist", self.lib)                  # an artist folder is not an album
        with self.assertRaises(replace.ReplaceError):
            replace.resolve_folder(self.tmp.name, self.lib)
        with self.assertRaises(replace.ReplaceError):
            replace.resolve_folder("Artist/Nope", self.lib)

    def test_swap_moves_the_old_folder_to_the_trash_and_puts_the_new_one_in_its_place(self):
        st = self.staging()
        moved, carried = replace.swap_in(self.old, st, self.lib, trash=self.trash)
        self.assertEqual(sorted(p.name for p in self.old.iterdir()), ["01 Song.flac", "02 Song.flac", "03 Song.flac", "booklet.pdf", "cover.jpg"])
        self.assertEqual((self.old / "cover.jpg").read_bytes(), b"new cover", "the new cover wins")
        self.assertEqual(carried, ["booklet.pdf"])
        self.assertTrue(str(moved).startswith(str(self.trash)))
        self.assertEqual(sorted(p.name for p in moved.iterdir() if p.suffix == ".mp3"), ["01 Song.mp3", "02 Song.mp3", "03 Song.mp3"])
        self.assertFalse(st.exists())

    def test_the_old_cover_is_kept_when_the_new_album_has_none(self):
        replace.swap_in(self.old, self.staging(cover=False), self.lib, trash=self.trash)
        self.assertEqual((self.old / "cover.jpg").read_bytes(), b"old cover")

    def test_an_incomplete_staging_folder_changes_nothing(self):
        st = self.staging(marker=True)
        with self.assertRaises(replace.ReplaceError):
            replace.swap_in(self.old, st, self.lib, trash=self.trash)
        self.assertEqual(len(list(self.old.glob("*.mp3"))), 3)
        self.assertEqual(list(self.trash.iterdir()), [])

    def test_if_the_new_folder_cannot_move_into_place_the_original_is_put_back(self):
        st = self.staging()
        with mock.patch("iaflac.replace.os.rename", side_effect=OSError("boom")):
            with self.assertRaises(replace.ReplaceError):
                replace.swap_in(self.old, st, self.lib, trash=self.trash)
        self.assertEqual(len(list(self.old.glob("*.mp3"))), 3, "the original is back")
        self.assertTrue(st.exists())

    def test_without_a_usable_trash_the_old_folder_goes_to_a_hidden_folder_in_the_library(self):
        dest = replace.trash_destination(self.old, self.lib, trash=self.trash / "missing")
        self.assertEqual(dest.parent, self.lib / replace.FALLBACK_TRASH / "Artist")


@unittest.skipUnless(HAVE_FLAC and HAVE_FFMPEG, "flac and ffmpeg required")
class ReplaceFlow(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.items = [make_item("ta-log", "Test Artist - Test Album", "Test Artist", tracks=4, rip=True, downloads=50, seed=1)]

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        self.out, self.lib, self.trash = str(root / "runs"), root / "1Music", root / "trash"
        self.trash.mkdir()
        self.old = self.lib / "Test Artist" / "Test Album"
        self.old.mkdir(parents=True)
        for i in range(1, 6):
            (self.old / f"{i:02d} Old Song {i}.mp3").write_bytes(b"mp3")
        (self.old / "booklet.pdf").write_bytes(b"pdf")
        rows = sorted(DATA + MP3_ROWS, key=lambda r: r[0])
        self.xlsx = build_workbook(self.lib / "1Music_Library_Inventory.xlsx", rows=rows)
        self.original = self.xlsx.read_bytes()
        self.client = FakeIA(self.items)
        self.env = mock.patch.dict(os.environ, {"IAFLAC_CONFIG": str(root / "config.json"), "IAFLAC_TRASH": str(self.trash)})
        self.env.start()
        for k in ("IAFLAC_LIBRARY", "IAFLAC_INVENTORY"):
            os.environ.pop(k, None)
        self.client_kw = {}

    def tearDown(self):
        self.env.stop()
        self.tmp.cleanup()

    def args(self, *items, **kw):
        base = dict(items=list(items or ["Test Artist - Test Album"]), file=None, top=3, include_live=False, free_only=False,
                    deep="none", no_spectral=True, max_deep_mb=1500, out=self.out, rip_only=False, max_candidates=10,
                    search_only=False, select=None, library=str(self.lib), inventory=None, no_inventory=False, save=None,
                    yes=False, art="cover", rip_files="skip", track_names="library", replace=False, scan=False)
        return Namespace(**{**base, **kw})

    def songs(self):
        return [r for r in load_workbook(self.xlsx)["Library"].iter_rows(min_row=2, max_col=5, values_only=True) if r[0] == "Test Artist"]

    def trashed(self):
        return [p for p in self.trash.iterdir()]

    # ---- through find: an existing MP3 album is noticed and offered for replacement -----------------

    def test_find_notices_the_mp3_album_and_replaces_it_only_after_a_yes(self):
        code, out, err, asked = run_cli(cli.cmd_find, self.args(), self.client, answers=["1", "y"], interactive=True)
        self.assertEqual(code, 0, err)
        self.assertIn("REPLACES your album (5 MP3)", asked[1]["shown"])
        self.assertIn("Replace 1 of your albums", asked[1]["prompt"])
        self.assertIn("Trash", asked[1]["prompt"])
        self.assertEqual(sorted(p.name for p in self.old.glob("*.flac")), [f"0{i} Track {i}.flac" for i in range(1, 5)])
        self.assertEqual(list(self.old.glob("*.mp3")), [])
        self.assertTrue((self.old / "booklet.pdf").exists(), "non-audio files are kept")
        trashed = self.trashed()
        self.assertEqual(len(trashed), 1)
        self.assertEqual(len(list(trashed[0].glob("*.mp3"))), 5, "the old files are in the trash, not deleted")
        self.assertEqual([r[4] for r in self.songs()], ["FLAC"] * 4)
        self.assertEqual(inv.Inventory(self.xlsx).consistency_problems(), [])
        self.assertEqual([p.name for p in self.old.parent.iterdir()], ["Test Album"], "no staging folder left behind")

    def test_the_plan_warns_when_the_track_counts_differ(self):
        _, _, _, asked = run_cli(cli.cmd_find, self.args(), self.client, answers=["1", "n"], interactive=True)
        self.assertIn("the old album has 5 track(s) and the new one 4", asked[1]["shown"])

    def test_declining_leaves_everything_exactly_as_it_was(self):
        _, out, _, _ = run_cli(cli.cmd_find, self.args(), self.client, answers=["1", "n"], interactive=True)
        self.assertEqual(len(list(self.old.glob("*.mp3"))), 5)
        self.assertEqual(self.trashed(), [])
        self.assertEqual(self.xlsx.read_bytes(), self.original)
        self.assertEqual([p.name for p in self.old.parent.iterdir()], ["Test Album"])

    def test_a_failed_download_leaves_the_original_untouched(self):
        client = FakeIA(self.items, corrupt_stream={"ta-log": "02 Track 2.flac"})
        code, out, err, _ = run_cli(cli.cmd_find, self.args(), client, answers=["1", "y"], interactive=True)
        self.assertEqual(code, 1)
        self.assertEqual(len(list(self.old.glob("*.mp3"))), 5)
        self.assertEqual(list(self.old.glob("*.flac")), [])
        self.assertEqual(self.trashed(), [])
        self.assertEqual(self.xlsx.read_bytes(), self.original)
        self.assertIn("your original album is untouched", out)

    def test_an_album_that_already_has_flac_is_never_replaced(self):
        (self.old / "01 Already.flac").write_bytes(b"f")
        code, out, err, asked = run_cli(cli.cmd_find, self.args(), self.client, answers=["1", "n"], interactive=True)
        self.assertNotIn("REPLACES", asked[1]["shown"] if len(asked) > 1 else out)

    def test_unattended_replacement_needs_the_replace_flag_as_well(self):
        code, out, _, _ = run_cli(cli.cmd_find, self.args(select="1", save="1", yes=True), self.client, interactive=False)
        self.assertIn("needs --replace", out)
        self.assertEqual(len(list(self.old.glob("*.mp3"))), 5)
        self.assertEqual(self.trashed(), [])
        code, out, err, _ = run_cli(cli.cmd_find, self.args(select="1", save="1", yes=True, replace=True), self.client, interactive=False)
        self.assertEqual(code, 0, err)
        self.assertEqual(list(self.old.glob("*.mp3")), [])
        self.assertEqual(len(self.trashed()), 1)

    def test_replacing_twice_is_harmless(self):
        run_cli(cli.cmd_find, self.args(select="1", save="1", yes=True, replace=True), self.client, interactive=False)
        first = self.xlsx.read_bytes()
        code, out, err, _ = run_cli(cli.cmd_find, self.args(select="1", save="1", yes=True, replace=True), self.client, interactive=False)
        self.assertNotIn("REPLACES", out, "the album now has FLAC, so it is no longer a replacement")
        self.assertIn("booklet.pdf", err, "the normal rule applies: it will not merge into a folder holding your own files")
        self.assertEqual(self.xlsx.read_bytes(), first)
        self.assertEqual(len(self.trashed()), 1, "nothing more was moved")
        self.assertEqual(len(list(self.old.glob("*.flac"))), 4)

    # ---- the replace command -------------------------------------------------------------------------

    def test_scan_lists_albums_that_have_audio_but_no_flac_and_changes_nothing(self):
        code, out, _, _ = run_cli(cli.cmd_replace, self.args(scan=True, items=[]), self.client)
        self.assertEqual(code, 0)
        self.assertIn("Test Artist/Test Album   (5 MP3)", out)
        self.assertEqual(len(list(self.old.glob("*.mp3"))), 5)

    def test_replace_takes_an_album_folder_and_shows_the_same_plan(self):
        code, out, err, asked = run_cli(cli.cmd_replace, self.args("Test Artist/Test Album", items=["Test Artist/Test Album"]),
                                        self.client, answers=["1", "y"], interactive=True)
        self.assertEqual(code, 0, err)
        self.assertEqual(list(self.old.glob("*.mp3")), [])
        self.assertEqual(len(list(self.old.glob("*.flac"))), 4)

    def test_replace_skips_a_folder_that_already_has_flac(self):
        (self.old / "01 Already.flac").write_bytes(b"f")
        code, out, err, _ = run_cli(cli.cmd_replace, self.args(items=["Test Artist/Test Album"]), self.client, interactive=False)
        self.assertEqual(code, 2)
        self.assertIn("already has FLAC", err)

    def test_replace_without_anything_named_explains_itself(self):
        code, _, err, _ = run_cli(cli.cmd_replace, self.args(items=[]), self.client)
        self.assertEqual(code, 2)
        self.assertIn("--scan", err)


if __name__ == "__main__":
    unittest.main()
