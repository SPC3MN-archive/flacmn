"""Add downloaded albums to the music inventory workbook (`Library` and `Summary` sheets).

The workbook's conventions were reverse-engineered from the real file and checked against it:

  Library  one row per track: Artist | Album | Song | File Type | Format
           artists form contiguous blocks sorted case-SENSITIVELY (Python's default order: "BROCKHAMPTON"
           before "Baby Keem"); albums inside a block are in no particular order; Album is the folder name;
           Song is the file name minus track number, minus a leading "Artist - "; File Type for FLAC is
           "FLAC (16-bit/44.1kHz)". Rows are banded by row parity (even rows shaded).
  Summary  one static row per artist, sorted case-INSENSITIVELY: Artist | # Albums | # Songs | Formats Present
           ("MP3 (19), FLAC (7)": biggest first), then a blank row, then TOTAL (albums, songs).

Rewriting the affected rows is safe here (no formulas, tables, charts or validations), but it is the user's data,
so every write is: lock-file check -> temp file -> re-read and verify -> timestamped backup -> atomic replace.
"""

from __future__ import annotations

import os
import posixpath
import re
import shutil
import time
import zipfile
from xml.etree import ElementTree as ET
from collections import Counter, OrderedDict
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from .flac_probe import probe_header
from .model import natural_key

LIBRARY, SUMMARY = "Library", "Summary"
LIBRARY_HEADER = ["Artist", "Album", "Song", "File Type", "Format"]
SUMMARY_HEADER = ["Artist", "# Albums", "# Songs", "Formats Present"]

LEADING_NUMBER = re.compile(r"^\s*(?:\d{1,2}[-.]\d{1,3}|\d{1,3})(?=[\s._\-)])[\s._\-)]*")
TRAILING_BRACKETS = re.compile(r"(?:\s*\[[^\]]*\])+\s*$")
UNSUPPORTED_AUDIO = (".mp3", ".m4a", ".wav", ".aiff", ".aif", ".ogg", ".opus", ".ape", ".wv", ".alac")


class InventoryError(RuntimeError):
    pass


# ---- deriving rows from files -------------------------------------------------------------------

def clean_song(filename: str, artist: str) -> str:
    """The inventory's Song: file name without track number, without a leading 'Artist - ', without trailing
    [bracketed notes] or the underscore a forbidden character (?) leaves behind. Never empty."""
    stem = os.path.splitext(os.path.basename(filename))[0]
    song = LEADING_NUMBER.sub("", stem)
    prefix = f"{artist} - "
    if song.casefold().startswith(prefix.casefold()):
        song = song[len(prefix):]
    song = TRAILING_BRACKETS.sub("", song).rstrip("_ ").strip()
    return song or stem.strip()


def flac_file_type(bits_per_sample: int, sample_rate: int) -> str:
    return f"FLAC ({bits_per_sample}-bit/{sample_rate / 1000:.1f}kHz)"


@dataclass
class NewAlbum:
    artist: str
    album: str
    rows: list                      # [(artist, album, song, file_type, format)] in track order
    folder: Optional[Path] = None
    replaces: str = ""              # an album already in the inventory (same artist) whose rows this one takes over


def rows_for_entries(entries: list, artist: str, album: str) -> list:
    """entries: [(relative path, bits_per_sample, sample_rate)] -> inventory rows in natural track order."""
    ordered = sorted(entries, key=lambda e: natural_key(str(e[0])))
    return [(artist, album, clean_song(str(path), artist), flac_file_type(bits, rate), "FLAC")
            for path, bits, rate in ordered]


def rows_for_typed(entries: list, artist: str, album: str) -> list:
    """entries: [(relative path, file type, format)] -> inventory rows in natural track order."""
    ordered = sorted(entries, key=lambda e: natural_key(str(e[0])))
    return [(artist, album, clean_song(str(path), artist), file_type, fmt) for path, file_type, fmt in ordered]


def mp3_file_type(path: Path) -> str:
    """'MP3 (VBR 245kbps)', in the style of the workbook's 'MP3 (CBR 128kbps)' entries."""
    from mutagen.mp3 import MP3
    try:
        info = MP3(path).info
    except Exception as exc:
        raise InventoryError(f"{path} is not a readable MP3 ({exc})") from exc
    mode = str(getattr(info.bitrate_mode, "name", "") or "").upper()
    return f"MP3 ({'VBR' if mode in ('', 'UNKNOWN') else mode} {round((info.bitrate or 0) / 1000)}kbps)"


def album_from_alt(item, stats, artist: str, album: str) -> NewAlbum:
    """A preview of the rows saving an upload in another format would produce (lossless becomes FLAC; MP3 and AAC stay)."""
    from .altmedia import is_lossless
    fmt = "FLAC" if is_lossless(item.media, stats) else ("AAC" if item.media == "m4a" else "MP3")
    return NewAlbum(artist, album, rows_for_typed([(f.name, stats.file_type, fmt) for f in item.flacs], artist, album))


def album_from_item(item, probes: dict, artist: str, album: str) -> NewAlbum:
    """A preview of the rows a download would produce, from archive.org's file list and the FLAC headers already
    probed. Tracks whose header was not probed are assumed to match the majority format of the ones that were."""
    formats = Counter((p.streaminfo.bits_per_sample, p.streaminfo.sample_rate)
                      for p in probes.values() if p.ok and p.streaminfo)
    if not formats:
        raise InventoryError("no probed FLAC headers to build a preview from")
    default = formats.most_common(1)[0][0]
    entries = []
    for f in item.flacs:
        p = probes.get(f.name)
        bits, rate = (p.streaminfo.bits_per_sample, p.streaminfo.sample_rate) if p and p.ok and p.streaminfo else default
        entries.append((f.name, bits, rate))
    return NewAlbum(artist, album, rows_for_entries(entries, artist, album))


def album_from_folder(folder: Path, artist: str, album: str, allow_mp3: bool = False) -> NewAlbum:
    """Read an album folder on disk (FLAC only, or MP3 only when allow_mp3) and build its rows from the real file headers."""
    entries, other, mp3s = [], [], []
    for root, dirs, files in os.walk(folder):
        dirs[:] = [d for d in dirs if not d.startswith(".")]
        for name in files:
            if name.startswith("."):                   # macOS ._ sidecars on FAT32/exFAT, .DS_Store
                continue
            path = Path(root) / name
            low = name.lower()
            if low.endswith(".flac"):
                with open(path, "rb") as fh:
                    probe = probe_header(lambda start, n, fh=fh: (fh.seek(start), fh.read(n))[1])
                if not probe.ok or not probe.streaminfo:
                    raise InventoryError(f"{path} is not a valid FLAC ({probe.kind}: {probe.detail})")
                si = probe.streaminfo
                entries.append((path.relative_to(folder), si.bits_per_sample, si.sample_rate))
            elif allow_mp3 and low.endswith(".mp3"):
                mp3s.append((path.relative_to(folder), mp3_file_type(path), "MP3"))
            elif low.endswith(UNSUPPORTED_AUDIO):
                other.append(name)
    if mp3s and not entries and not other:
        return NewAlbum(artist, album, rows_for_typed(mp3s, artist, album), folder)
    if other or mp3s:
        other += [p[0].name for p in mp3s]
    if other:
        raise InventoryError(f"{folder} contains non-FLAC audio ({', '.join(sorted(other)[:3])}...); "
                             "only FLAC albums are supported")
    if not entries:
        raise InventoryError(f"no FLAC files found in {folder}")
    return NewAlbum(artist, album, rows_for_entries(entries, artist, album), folder)


# ---- summarising ------------------------------------------------------------------------------

def summarise(rows: list) -> "OrderedDict[str, tuple]":
    """{artist: (# albums, # songs, 'FMT (n), FMT (n)')}, the way the Summary sheet computes it."""
    by_artist: OrderedDict = OrderedDict()
    for artist, album, _song, _ft, fmt in rows:
        by_artist.setdefault(artist, []).append((album, fmt))
    out: OrderedDict = OrderedDict()
    for artist, items in by_artist.items():
        counts = Counter(fmt for _, fmt in items)
        formats = ", ".join(f"{k} ({v})" for k, v in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0])))
        out[artist] = (len({album for album, _ in items}), len(items), formats)
    return out


# ---- planning ---------------------------------------------------------------------------------

@dataclass
class Placement:
    album: NewAlbum                  # with the artist spelled as it will appear
    status: str                      # "new artist" | "existing artist" | "replaces existing rows" | "already in inventory"
    library_row: int = 0             # first Library sheet row the album will occupy
    summary_row: int = 0
    removed: list = field(default_factory=list)      # the old rows a replacement takes out

    @property
    def added(self) -> bool:
        return self.status != "already in inventory"


@dataclass
class Plan:
    placements: list
    rows_before: int
    rows_after: int
    totals_before: tuple             # (albums, songs)
    totals_after: tuple
    library_rows: list = field(default_factory=list)      # the full new Library data
    summary_rows: list = field(default_factory=list)      # [(artist, albums, songs, formats)] in sheet order
    changed_artists: set = field(default_factory=set)
    removed_rows: list = field(default_factory=list)      # rows taken out by replacements (everything else is kept)

    @property
    def changes(self) -> bool:
        return any(p.added for p in self.placements)


def album_rows(rows: list, artist: str, album: str) -> list:
    """Indexes of an album's rows (the artist exactly, the album case-insensitively)."""
    return [i for i, r in enumerate(rows) if r[0] == artist and str(r[1]).casefold() == album.casefold()]


def artist_insert_index(rows: list, artist: str) -> int:
    """Where a new artist block goes: before the first block whose name sorts after it (case-sensitively)."""
    for i, row in enumerate(rows):
        if row[0] > artist:
            return i
    return len(rows)


def end_of_block(rows: list, artist: str) -> int:
    last = max(i for i, r in enumerate(rows) if r[0] == artist)
    return last + 1


class Inventory:
    def __init__(self, path):
        self.path = Path(path).expanduser()
        if not self.path.exists():
            raise InventoryError(f"{self.path} does not exist")
        self.wb = load_workbook(self.path)
        for name in (LIBRARY, SUMMARY):
            if name not in self.wb.sheetnames:
                raise InventoryError(f"workbook has no {name!r} sheet (found {self.wb.sheetnames})")
        self.lib, self.sum = self.wb[LIBRARY], self.wb[SUMMARY]
        if [c.value for c in self.lib[1]][:5] != LIBRARY_HEADER or [c.value for c in self.sum[1]][:4] != SUMMARY_HEADER:
            raise InventoryError("sheet headers are not the expected Artist/Album/Song/File Type/Format and "
                                 "Artist/# Albums/# Songs/Formats Present")
        self.library_rows = [tuple(r) for r in self.lib.iter_rows(min_row=2, max_col=5, values_only=True)
                             if r[0] is not None]
        self.summary_rows, self.total_row = self._read_summary()
        self._check_banding()

    # -- reading ------------------------------------------------------------------------

    def _read_summary(self) -> tuple:
        artists, r = [], 2
        while self.sum.cell(r, 1).value not in (None, "TOTAL"):
            artists.append(tuple(self.sum.cell(r, c).value for c in range(1, 5)))
            r += 1
        blank, total_row = r, r + 1
        if (any(self.sum.cell(blank, c).value is not None for c in range(1, 5))
                or self.sum.cell(total_row, 1).value != "TOTAL"):
            raise InventoryError("Summary layout is not 'artist rows, one blank row, TOTAL'; not editing it")
        self._sum_blank_row, self._sum_total_row = blank, total_row
        return artists, tuple(self.sum.cell(total_row, c).value for c in range(1, 5))

    def _shaded(self, ws, row: int) -> bool:
        return ws.cell(row, 1).fill.fill_type == "solid"

    def _check_banding(self) -> None:
        """Rows alternate shading by parity in the real file; stop if someone changed that, rather than guess."""
        for ws, last in ((self.lib, 1 + len(self.library_rows)), (self.sum, 1 + len(self.summary_rows))):
            first_shaded = self._shaded(ws, 2)
            if not first_shaded or self._shaded(ws, 3):
                raise InventoryError(f"{ws.title}: expected rows to alternate shaded/plain starting with a shaded row 2")
            for r in range(2, last + 1):
                if self._shaded(ws, r) != (r % 2 == 0):
                    raise InventoryError(f"{ws.title}: row {r} breaks the alternating row shading; "
                                         "the sheet was edited by hand, so I won't guess how to extend it")

    def artists(self) -> list:
        return list(OrderedDict.fromkeys(r[0] for r in self.library_rows))

    def canonical_artist(self, artist: str) -> str:
        """Use the spelling already in the inventory ('Boards of Canada', not 'Boards Of Canada')."""
        for existing in self.artists():
            if existing.casefold() == artist.casefold():
                return existing
        return artist

    def consistency_problems(self) -> list:
        """Does the Summary agree with the Library? (It does in the real file: 0 differences.)"""
        computed = summarise(self.library_rows)
        problems = []
        seen = set()
        for artist, albums, songs, fmts in self.summary_rows:
            seen.add(artist)
            if computed.get(artist) != (albums, songs, fmts):
                problems.append(f"{artist}: Summary says {(albums, songs, fmts)}, Library gives {computed.get(artist)}")
        problems += [f"{a}: in Library but not in Summary" for a in computed if a not in seen]
        albums_total = sum(v[0] for v in computed.values())
        if self.total_row[1] != albums_total or self.total_row[2] != len(self.library_rows):
            problems.append(f"TOTAL says {self.total_row[1:3]}, Library gives {(albums_total, len(self.library_rows))}")
        return problems

    # -- planning -----------------------------------------------------------------------

    def plan(self, albums: list) -> Plan:
        rows = list(self.library_rows)
        albums_before = sum(v[0] for v in summarise(rows).values())
        placements, changed, removed_rows = [], set(), []
        for al in albums:
            artist = self.canonical_artist(al.artist)
            fixed = NewAlbum(artist, al.album, [(artist, al.album, s, ft, fm) for _a, _b, s, ft, fm in al.rows], al.folder)
            old_idx = album_rows(rows, artist, al.replaces) if al.replaces else []
            if old_idx:
                # Replace in place: the new rows go where the old album's rows were, nothing else moves.
                old = [rows[i] for i in old_idx]
                if old == fixed.rows:
                    placements.append(Placement(fixed, "already in inventory"))
                    continue
                if old_idx != list(range(old_idx[0], old_idx[0] + len(old_idx))):
                    raise InventoryError(f"the rows of {artist} / {al.replaces} are not together in the Library sheet; "
                                         "not replacing them automatically")
                rows[old_idx[0]:old_idx[-1] + 1] = fixed.rows
                changed.add(artist)
                removed_rows += old
                placements.append(Placement(fixed, "replaces existing rows", removed=old))
                continue
            if any(r[0] == artist and str(r[1]).casefold() == al.album.casefold() for r in rows):
                placements.append(Placement(fixed, "already in inventory"))
                continue
            exists = any(r[0] == artist for r in rows)
            idx = end_of_block(rows, artist) if exists else artist_insert_index(rows, artist)
            rows[idx:idx] = fixed.rows
            changed.add(artist)
            placements.append(Placement(fixed, "existing artist" if exists else "new artist"))
        # final positions, after every insertion
        for p in placements:
            if p.added:
                p.library_row = 2 + next(i for i, r in enumerate(rows) if r[0] == p.album.artist and r[1] == p.album.album)
        summary = self._new_summary(rows, changed)
        for p in placements:
            if p.added:
                p.summary_row = 2 + next(i for i, s in enumerate(summary) if s[0] == p.album.artist)
        computed = summarise(rows)
        return Plan(placements=placements, rows_before=len(self.library_rows), rows_after=len(rows),
                    totals_before=(albums_before, len(self.library_rows)),
                    totals_after=(sum(v[0] for v in computed.values()), len(rows)),
                    library_rows=rows, summary_rows=summary, changed_artists=changed, removed_rows=removed_rows)

    def _new_summary(self, rows: list, changed: set) -> list:
        """Existing Summary rows stay exactly as they are; only changed artists are recomputed, new ones inserted
        at their case-insensitive position."""
        computed = summarise(rows)
        by_artist = OrderedDict((r[0], r) for r in self.summary_rows)
        for artist in changed:
            by_artist[artist] = (artist, *computed[artist])
        return sorted(by_artist.values(), key=lambda r: str(r[0]).casefold())

    # -- writing ------------------------------------------------------------------------

    def apply(self, plan: Plan, backup: bool = True) -> dict:
        """Write the plan to the file. Returns {'backup': Path|None, 'written': bool}."""
        if not plan.changes:
            return {"backup": None, "written": False}
        for lock in (self.path.with_name("~$" + self.path.name), self.path.with_name(f".~lock.{self.path.name}#")):
            if lock.exists():
                raise InventoryError(f"{lock.name} exists: the workbook is open in a spreadsheet app "
                                     "(Excel or LibreOffice). Close it and try again.")

        if len(plan.summary_rows) < len(self.summary_rows):
            raise InventoryError("this change would remove an artist from the Summary, which is not supported; nothing was changed")
        wb = load_workbook(self.path)                  # a fresh copy: `self.wb` stays untouched
        lib, summ = wb[LIBRARY], wb[SUMMARY]
        lib_style = {0: [copy(lib.cell(2, c)._style) for c in range(1, 6)], 1: [copy(lib.cell(3, c)._style) for c in range(1, 6)]}
        sum_style = {0: [copy(summ.cell(2, c)._style) for c in range(1, 5)], 1: [copy(summ.cell(3, c)._style) for c in range(1, 5)]}
        blank_style = [copy(summ.cell(self._sum_blank_row, c)._style) for c in range(1, 5)]
        total_style = [copy(summ.cell(self._sum_total_row, c)._style) for c in range(1, 5)]

        for i, values in enumerate(plan.library_rows):
            r = i + 2
            for j, value in enumerate(values):
                cell = lib.cell(r, j + 1)
                cell.value = value
                cell._style = copy(lib_style[r % 2][j])
        stale = lib.max_row - (1 + len(plan.library_rows))
        if stale > 0:                                  # a replacement can shrink the sheet: drop the leftover rows
            lib.delete_rows(2 + len(plan.library_rows), stale)
        lib.auto_filter.ref = f"A1:E{1 + len(plan.library_rows)}"

        for i, values in enumerate(plan.summary_rows):
            r = i + 2
            for j, value in enumerate(values):
                cell = summ.cell(r, j + 1)
                cell.value = value
                cell._style = copy(sum_style[r % 2][j])
        last = 1 + len(plan.summary_rows)
        for j in range(4):
            cell = summ.cell(last + 1, j + 1)
            cell.value = None
            cell._style = copy(blank_style[j])
        total = ("TOTAL", sum(r[1] for r in plan.summary_rows), sum(r[2] for r in plan.summary_rows), None)
        for j, value in enumerate(total):
            cell = summ.cell(last + 2, j + 1)
            cell.value = value
            cell._style = copy(total_style[j])
        summ.auto_filter.ref = f"A1:D{last}"

        tmp = self.path.with_name(f".{self.path.stem}.iaflac-tmp{self.path.suffix}")
        wb.save(tmp)
        try:
            problems = verify_written(tmp, plan, reference=self.path)
            if problems:
                raise InventoryError("the edited workbook failed verification, so nothing was changed: " + "; ".join(problems[:5]))
            backup_path = None
            if backup:
                backup_path = self.path.with_name(f"{self.path.stem}.backup-{time.strftime('%Y%m%d-%H%M%S')}{self.path.suffix}")
                shutil.copy2(self.path, backup_path)
            shutil.copymode(self.path, tmp)             # keep the file's permissions across the replace
            os.replace(tmp, self.path)
        finally:
            if tmp.exists():
                tmp.unlink()
        return {"backup": backup_path, "written": True}


# ---- verification -----------------------------------------------------------------------------

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def _col_index(ref: str) -> int:
    letters = re.match(r"[A-Z]+", ref).group(0)
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def ooxml_problems(path: Path) -> list:
    """Structural checks on the .xlsx package that do not rely on openpyxl (the library that wrote it):
    valid zip, well-formed XML, content types and relationships resolve, rows and cells are in order with no
    duplicates, style indexes exist, and the filter ranges agree. These are the things Excel's 'repair' prompt
    is triggered by."""
    problems = []
    try:
        z = zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        return [f"not a valid zip: {exc}"]
    with z:
        if z.testzip() is not None:
            return [f"corrupt zip member {z.testzip()}"]
        names = set(z.namelist())
        docs = {}
        for name in names:
            if name.endswith((".xml", ".rels")):
                try:
                    docs[name] = ET.fromstring(z.read(name))
                except ET.ParseError as exc:
                    problems.append(f"{name} is not well-formed XML: {exc}")
        if problems:
            return problems
        ct = docs.get("[Content_Types].xml")
        if ct is None:
            return ["[Content_Types].xml is missing"]
        defaults = {e.get("Extension") for e in ct if e.tag.endswith("Default")}
        overrides = {e.get("PartName").lstrip("/") for e in ct if e.tag.endswith("Override")}
        for name in names:
            if name != "[Content_Types].xml" and name.rsplit(".", 1)[-1] not in defaults and name not in overrides:
                problems.append(f"{name} has no content type")
        for rels_name, rels in docs.items():
            if not rels_name.endswith(".rels"):
                continue
            base = posixpath.dirname(posixpath.dirname(rels_name))
            for rel in rels:
                if rel.get("TargetMode") == "External":
                    continue
                target = rel.get("Target")
                resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                if resolved not in names:
                    problems.append(f"{rels_name}: relationship target {resolved} is missing")
        styles = docs.get("xl/styles.xml")
        n_xfs = len(styles.find("m:cellXfs", NS)) if styles is not None and styles.find("m:cellXfs", NS) is not None else 0
        wb = docs.get("xl/workbook.xml")
        filter_names = {}
        if wb is not None:
            for dn in wb.iterfind(".//m:definedName", NS):
                if dn.get("name") == "_xlnm._FilterDatabase":
                    filter_names[int(dn.get("localSheetId"))] = dn.text.replace("'", "").replace("$", "")
        for idx, name in enumerate(sorted(n for n in names if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n))):
            sheet = docs[name]
            last_row, max_row, max_col = 0, 0, 0
            for row in sheet.iterfind(".//m:sheetData/m:row", NS):
                r = int(row.get("r"))
                if r <= last_row:
                    problems.append(f"{name}: row {r} is not after row {last_row}")
                last_row, max_row, last_col = r, max(max_row, r), 0
                for c in row.iterfind("m:c", NS):
                    ref = c.get("r")
                    if not re.fullmatch(r"[A-Z]+%d" % r, ref or ""):
                        problems.append(f"{name}: cell {ref} sits in row {r}")
                        continue
                    col = _col_index(ref)
                    if col <= last_col:
                        problems.append(f"{name}: cell {ref} is out of order or duplicated")
                    last_col, max_col = col, max(max_col, col)
                    if c.get("s") is not None and int(c.get("s")) >= n_xfs:
                        problems.append(f"{name}: cell {ref} uses style {c.get('s')} but only {n_xfs} exist")
            dim = sheet.find("m:dimension", NS)
            if dim is not None and dim.get("ref") and ":" in dim.get("ref"):
                end = dim.get("ref").split(":")[1]
                if _col_index(end) < max_col or int(re.search(r"\d+", end).group(0)) < max_row:
                    problems.append(f"{name}: dimension {dim.get('ref')} does not cover the used cells")
            af = sheet.find("m:autoFilter", NS)
            if af is not None and idx in filter_names and filter_names[idx].split("!")[-1] != af.get("ref"):
                problems.append(f"{name}: autoFilter {af.get('ref')} disagrees with its hidden filter name {filter_names[idx]}")
    return problems


def _sheet_facts(ws) -> dict:
    pane = ws.sheet_view.pane
    return {"widths": {k: d.width for k, d in ws.column_dimensions.items() if d.width},
            "zoom": ws.sheet_view.zoomScale, "pane": (pane.xSplit, pane.ySplit, pane.state) if pane else None,
            "header_height": ws.row_dimensions[1].height,
            "header": [(c.value, c.font.name, c.font.sz, c.font.b, c.fill.fill_type) for c in ws[1]]}


def verify_written(new_path: Path, plan: Plan, reference: Path) -> list:
    """Re-read the written file and compare it with what we meant to write and with the original."""
    problems = ooxml_problems(new_path)
    if problems:
        return problems
    new, old = load_workbook(new_path), load_workbook(reference)
    if new.sheetnames != old.sheetnames:
        return [f"sheets changed: {old.sheetnames} -> {new.sheetnames}"]
    lib, summ = new[LIBRARY], new[SUMMARY]
    got = [tuple(r) for r in lib.iter_rows(min_row=2, max_col=5, values_only=True)]
    if got != list(plan.library_rows):
        problems.append("Library rows differ from the plan")
    if lib.max_row != 1 + len(plan.library_rows):
        problems.append(f"Library has {lib.max_row} rows, expected {1 + len(plan.library_rows)}")
    for r in range(2, 2 + len(plan.library_rows)):
        if any((lib.cell(r, c).fill.fill_type == "solid") != (r % 2 == 0) for c in range(1, 6)):
            problems.append(f"Library row {r} has the wrong banding")
            break
    got_sum = [tuple(r) for r in summ.iter_rows(min_row=2, max_row=1 + len(plan.summary_rows), max_col=4, values_only=True)]
    if got_sum != [tuple(r) for r in plan.summary_rows]:
        problems.append("Summary rows differ from the plan")
    last = 1 + len(plan.summary_rows)
    if [summ.cell(last + 2, c).value for c in range(1, 4)] != ["TOTAL", sum(r[1] for r in plan.summary_rows), sum(r[2] for r in plan.summary_rows)]:
        problems.append("TOTAL row is wrong")
    if any(summ.cell(last + 1, c).value is not None for c in range(1, 5)):
        problems.append("the blank row before TOTAL is not blank")
    for r in range(2, last + 1):
        if any((summ.cell(r, c).fill.fill_type == "solid") != (r % 2 == 0) for c in range(1, 5)):
            problems.append(f"Summary row {r} has the wrong banding")
            break
    if lib.auto_filter.ref != f"A1:E{1 + len(plan.library_rows)}" or summ.auto_filter.ref != f"A1:D{last}":
        problems.append(f"autofilter ranges are {lib.auto_filter.ref} / {summ.auto_filter.ref}")
    for name in (LIBRARY, SUMMARY):
        if _sheet_facts(new[name]) != _sheet_facts(old[name]):
            problems.append(f"{name}: column widths, zoom, frozen panes or header formatting changed")
    # every original Library row is still there, in order
    gone = list(plan.removed_rows)                     # only rows a replacement was asked to take out may disappear
    old_rows = []
    for r in old[LIBRARY].iter_rows(min_row=2, max_col=5, values_only=True):
        if r[0] is None:
            continue
        if tuple(r) in gone:
            gone.remove(tuple(r))
            continue
        old_rows.append(tuple(r))
    if gone:
        problems.append("a row to be replaced was not in the original workbook")
    it = iter(got)
    if not all(any(row == g for g in it) for row in old_rows):
        problems.append("an existing Library row was lost or reordered")
    return problems
