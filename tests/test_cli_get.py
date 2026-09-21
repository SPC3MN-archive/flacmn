"""The whole 'download an album into my library and add it to my inventory' flow, through the real CLI handlers,
against a fake archive that serves real (tiny) FLACs and a synthetic inventory workbook."""

import contextlib
import io
import tempfile
import unittest
from argparse import Namespace
from pathlib import Path

from openpyxl import load_workbook

from iaflac import cli
from tests.helpers import HAVE_FFMPEG, HAVE_FLAC
from tests.test_inventory import DATA, build_workbook
from tests.test_pipeline import FakeIA, make_item

PORTABLE = dict(dest=None, no_art=False, art="cover", rip_files="keep", as_names=None)


def run(func, args, client):
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        code = func(args, client)
    return code, out.getvalue(), err.getvalue()


@unittest.skipUnless(HAVE_FLAC and HAVE_FFMPEG, "flac and ffmpeg required")
class LibraryAndInventoryFlow(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.items = [
            make_item("mv-1", "The Mars Volta - Frances the Mute (075021039773 - 2005-02-21)", "The Mars Volta", tracks=4,
                      rip=True, seed=1, album_tag="Frances the Mute"),
            make_item("boc-1", "Boards Of Canada - Geogaddi (warpcd101)", "", tracks=3, rip=False, seed=2,
                      artist_tag="", album_tag=""),
        ]

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        self.lib = root / "1Music"
        (self.lib / "Radiohead" / "OK Computer").mkdir(parents=True)
        (self.lib / "Radiohead" / "OK Computer" / "01 Airbag.flac").write_bytes(b"mine")
        self.xlsx = build_workbook(self.lib / "1Music_Library_Inventory.xlsx")
        self.original = self.xlsx.read_bytes()
        self.client = FakeIA(self.items)

    def tearDown(self):
        self.tmp.cleanup()

    def args(self, ids, yes=False, inventory=True, **kw):
        return Namespace(identifiers=ids, library=str(self.lib), inventory=str(self.xlsx) if inventory else None,
                         yes=yes, **{**PORTABLE, **kw})

    def songs(self):
        ws = load_workbook(self.xlsx)["Library"]
        return [tuple(r) for r in ws.iter_rows(min_row=2, max_col=5, values_only=True)]

    def test_dry_run_shows_the_plan_and_the_rows_and_changes_nothing(self):
        code, out, _ = run(cli.cmd_get, self.args(["mv-1"]), self.client)
        self.assertEqual(code, 0)
        self.assertIn(str(self.lib / "The Mars Volta" / "Frances the Mute"), out)
        self.assertIn("new artist", out)
        self.assertIn("The Mars Volta | Frances the Mute | Track 1 | FLAC (16-bit/44.1kHz) | FLAC", out)
        self.assertIn("Nothing was downloaded and nothing was changed", out)
        self.assertEqual(self.xlsx.read_bytes(), self.original)
        self.assertFalse((self.lib / "The Mars Volta").exists())
        self.assertEqual(self.client.bytes_read, 0)

    def test_download_then_inventory_update_with_backup(self):
        code, out, err = run(cli.cmd_get, self.args(["mv-1"], yes=True), self.client)
        self.assertEqual(code, 0, out + err)
        folder = self.lib / "The Mars Volta" / "Frances the Mute"
        self.assertEqual(sorted(p.name for p in folder.iterdir() if p.suffix == ".flac"), [f"0{i} Track {i}.flac" for i in range(1, 5)])
        self.assertTrue((folder / "rip.log").exists() and (folder / "album.cue").exists())
        rows = self.songs()
        new = [r for r in rows if r[0] == "The Mars Volta"]
        self.assertEqual([r[2] for r in new], ["Track 1", "Track 2", "Track 3", "Track 4"])
        artists = list(dict.fromkeys(r[0] for r in rows))
        self.assertEqual(artists.index("The Mars Volta"), artists.index("The Marias") + 1)
        self.assertEqual(len(rows), len(DATA) + 4)
        backups = list(self.lib.glob("1Music_Library_Inventory.backup-*.xlsx"))
        self.assertEqual(len(backups), 1)
        self.assertEqual(backups[0].read_bytes(), self.original)
        self.assertIn("backup of the previous version", out)
        self.assertEqual((self.lib / "Radiohead" / "OK Computer" / "01 Airbag.flac").read_bytes(), b"mine", "existing music untouched")

    def test_names_are_guessed_shown_and_the_existing_spelling_of_the_artist_is_reused(self):
        code, out, _ = run(cli.cmd_get, self.args(["boc-1"], yes=True), self.client)
        self.assertEqual(code, 0, out)
        self.assertIn("GUESSED from the title", out)
        rows = [r for r in self.songs() if r[1] == "Geogaddi"]
        self.assertEqual({r[0] for r in rows}, {"Boards of Canada"}, "the inventory's spelling, not the archive's 'Of'")
        self.assertTrue((self.lib / "Boards of Canada" / "Geogaddi" / "01 Track 1.flac").exists())

    def test_explicit_names_override_the_guess(self):
        code, out, _ = run(cli.cmd_get, self.args(["boc-1"], as_names=["boc-1=Boards of Canada | Geogaddi (Test)"]), self.client)
        self.assertEqual(code, 0)
        self.assertIn("given with --as", out)
        self.assertIn("Geogaddi (Test)", out)

    def test_rerunning_changes_nothing_and_makes_no_second_backup(self):
        run(cli.cmd_get, self.args(["mv-1"], yes=True), self.client)
        after_first = self.xlsx.read_bytes()
        self.client.bytes_read = 0
        code, out, _ = run(cli.cmd_get, self.args(["mv-1"], yes=True), self.client)
        self.assertEqual(code, 0, out)
        self.assertIn("already had", out)
        self.assertIn("ALREADY in the inventory", out)
        self.assertEqual(self.xlsx.read_bytes(), after_first)
        self.assertEqual(len(list(self.lib.glob("*.backup-*"))), 1)

    def test_a_corrupt_download_is_not_added_to_the_inventory(self):
        client = FakeIA(self.items, corrupt_stream={"mv-1": "02 Track 2.flac"})
        code, out, _ = run(cli.cmd_get, self.args(["mv-1"], yes=True), client)
        self.assertEqual(code, 1)
        self.assertIn("FAILED", out)
        self.assertIn("NOT adding this album to the inventory", out)
        self.assertEqual(self.xlsx.read_bytes(), self.original)
        self.assertFalse((self.lib / "The Mars Volta" / "Frances the Mute" / "02 Track 2.flac").exists())
        self.assertTrue((self.lib / "The Mars Volta" / "Frances the Mute" / ".iaflac-incomplete").exists())

    def test_excel_having_the_workbook_open_means_downloaded_but_not_recorded_and_says_how_to_fix_it(self):
        (self.lib / "~$1Music_Library_Inventory.xlsx").write_bytes(b"lock")
        code, out, err = run(cli.cmd_get, self.args(["mv-1"], yes=True), self.client)
        self.assertEqual(code, 1)
        self.assertIn("NOT updated", err)
        self.assertIn("inventory add", err)
        self.assertTrue((self.lib / "The Mars Volta" / "Frances the Mute" / "01 Track 1.flac").exists())
        self.assertEqual(self.xlsx.read_bytes(), self.original)

    def test_an_existing_album_folder_of_yours_is_never_merged_into(self):
        mine = self.lib / "The Mars Volta" / "Frances the Mute"
        mine.mkdir(parents=True)
        (mine / "mine.flac").write_bytes(b"precious")
        code, _, err = run(cli.cmd_get, self.args(["mv-1"], yes=True), self.client)
        self.assertEqual(code, 2)
        self.assertIn("not merging", err)
        self.assertEqual([p.name for p in mine.iterdir()], ["mine.flac"])

    def test_inventory_needs_a_library(self):
        code, _, err = run(cli.cmd_get, Namespace(identifiers=["x"], library=None, dest="/tmp/x", inventory=str(self.xlsx),
                                                  yes=False, **{k: v for k, v in PORTABLE.items() if k != "dest"}), self.client)
        self.assertEqual(code, 2)
        self.assertIn("--inventory needs --library", err)

    def test_exactly_one_layout_must_be_chosen(self):
        for lib, dest in ((None, None), ("/a", "/b")):
            code, _, err = run(cli.cmd_get, Namespace(identifiers=["x"], library=lib, dest=dest, inventory=None, yes=False,
                                                      **{k: v for k, v in PORTABLE.items() if k != "dest"}), self.client)
            self.assertEqual(code, 2)


@unittest.skipUnless(HAVE_FLAC and HAVE_FFMPEG, "flac and ffmpeg required")
class InventoryCommands(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.xlsx = build_workbook(self.root / "inv.xlsx")
        self.original = self.xlsx.read_bytes()
        self.folder = self.root / "Miles Davis" / "Kind of Blue"
        self.folder.mkdir(parents=True)
        for i, name in enumerate(("01 - Miles Davis - So What", "02 - Miles Davis - Freddie Freeloader"), 1):
            (self.folder / f"{name}.flac").write_bytes(make_item("x", "t", "c", tracks=1, rip=False).files["01 Track 1.flac"])

    def tearDown(self):
        self.tmp.cleanup()

    def ns(self, action, **kw):
        return Namespace(action=action, xlsx=str(self.xlsx), folders=kw.pop("folders", []), artist=None, album=None,
                         yes=kw.pop("yes", False), **kw)

    def test_check_reports_a_consistent_workbook(self):
        code, out, _ = run(cli.cmd_inventory, self.ns("check"), None)
        self.assertEqual(code, 0)
        self.assertIn("Summary agrees with Library", out)

    def test_check_flags_a_disagreement(self):
        wb = load_workbook(self.xlsx)
        wb["Summary"]["C2"] = 999
        wb.save(self.xlsx)
        code, out, _ = run(cli.cmd_inventory, self.ns("check"), None)
        self.assertEqual(code, 1)
        self.assertIn("DISAGREE", out)

    def test_add_is_a_dry_run_by_default_and_names_come_from_the_folder_path(self):
        code, out, _ = run(cli.cmd_inventory, self.ns("add", folders=[str(self.folder)]), None)
        self.assertEqual(code, 0)
        self.assertIn("new artist: Miles Davis / Kind of Blue", out)
        self.assertIn("Miles Davis | Kind of Blue | So What", out)
        self.assertIn("Dry run", out)
        self.assertEqual(self.xlsx.read_bytes(), self.original)

    def test_add_with_yes_writes_and_backs_up(self):
        code, out, _ = run(cli.cmd_inventory, self.ns("add", folders=[str(self.folder)], yes=True), None)
        self.assertEqual(code, 0, out)
        rows = [tuple(r) for r in load_workbook(self.xlsx)["Library"].iter_rows(min_row=2, max_col=5, values_only=True)]
        self.assertEqual([r[2] for r in rows if r[0] == "Miles Davis"], ["So What", "Freddie Freeloader"])
        self.assertEqual(len(list(self.root.glob("inv.backup-*.xlsx"))), 1)

    def test_add_needs_folders_and_a_single_folder_for_name_overrides(self):
        self.assertEqual(run(cli.cmd_inventory, self.ns("add"), None)[0], 2)
        self.assertEqual(run(cli.cmd_inventory, Namespace(action="add", xlsx=str(self.xlsx), folders=["a", "b"], artist="X", album=None, yes=False), None)[0], 2)

    def test_a_missing_workbook_is_a_clear_error(self):
        code, _, err = run(cli.cmd_inventory, Namespace(action="check", xlsx=str(self.root / "nope.xlsx"), folders=[], artist=None, album=None, yes=False), None)
        self.assertEqual(code, 2)
        self.assertIn("does not exist", err)


if __name__ == "__main__":
    unittest.main()
