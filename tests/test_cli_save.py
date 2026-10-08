import json
"""find -> choose what to test -> choose what to save -> confirm -> saved in the library's own layout, through the
real CLI handlers against a fake archive and a synthetic library and inventory."""

import hashlib
import os
import tempfile
import unittest
from argparse import Namespace
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from iaflac import cli, download
from iaflac.model import FileInfo, ItemInfo
from tests.helpers import HAVE_FFMPEG, HAVE_FLAC
from tests.test_cli_find import file_reads, run_cli
from tests.test_inventory import build_workbook
from tests.test_pipeline import FakeIA, make_item


def info(name):
    return FileInfo(name=name, format="Flac", source="original", size=10, md5="0" * 32, sha1="0" * 40)


class LibraryTrackNames(unittest.TestCase):
    def name(self, raw, artist="Miles Davis"):
        return download.library_track_name(Path(raw), artist).as_posix()

    def test_archive_names_become_the_librarys_NN_Title(self):
        self.assertEqual(self.name("01 - Miles Davis - So What.flac"), "01 So What.flac")
        self.assertEqual(self.name("5. Miles Davis - Flamenco Sketches.flac"), "05 Flamenco Sketches.flac")
        self.assertEqual(self.name("03. Blue in Green.FLAC"), "03 Blue in Green.flac")
        self.assertEqual(self.name("7_All Blues.flac"), "07 All Blues.flac")

    def test_disc_track_numbers_and_unnumbered_names_are_left_alone(self):
        self.assertEqual(self.name("1-05 Goodbye.flac"), "1-05 Goodbye.flac")
        self.assertEqual(self.name("Frankie's First Affair.flac"), "Frankie's First Affair.flac")
        self.assertEqual(self.name("1999.flac"), "1999.flac")

    def test_a_title_that_is_only_a_number_prefix_is_not_emptied(self):
        self.assertEqual(self.name("07 - .flac"), "07 - .flac")

    def test_subfolders_are_kept(self):
        self.assertEqual(self.name("CD 2/04 - Miles Davis - Song.flac"), "CD 2/04 Song.flac")

    def plan(self, names, **kw):
        item = ItemInfo(identifier="x", title="Miles Davis - Album", files=[info(n) for n in names])
        return download.plan_library(item, Path("/tmp/iaflac-lib"), "Miles Davis", "Album", rip="skip", art="none", **kw)

    def rels(self, res):
        return sorted(str(pf.path.relative_to(res.folder.resolve())) for pf in res.planned)

    def test_plan_uses_the_library_style_when_asked_and_archive_names_otherwise(self):
        names = ["01 - Miles Davis - So What.flac", "02 - Miles Davis - Blue in Green.flac"]
        self.assertEqual(self.rels(self.plan(names, track_names="library")), ["01 So What.flac", "02 Blue in Green.flac"])
        self.assertEqual(self.rels(self.plan(names, track_names="keep")), sorted(names))

    def test_if_renaming_would_make_two_tracks_collide_every_archive_name_is_kept(self):
        names = ["01 - Miles Davis - Song.flac", "01. Song.flac"]
        self.assertEqual(self.rels(self.plan(names, track_names="library")), sorted(names))


class SaveHelpers(unittest.TestCase):
    def test_inventory_is_taken_from_the_library_folder_only_when_unambiguous(self):
        with tempfile.TemporaryDirectory() as tmp:
            lib = Path(tmp)
            args = Namespace(inventory=None, no_inventory=False)
            with mock.patch.dict(os.environ, {}, clear=False):
                os.environ.pop("IAFLAC_INVENTORY", None)
                self.assertEqual(cli.inventory_from(args, lib)[0], None)
                (lib / "1Music_Library_Inventory.xlsx").write_bytes(b"x")
                (lib / "1Music_Library_Inventory.backup-20260101-000000.xlsx").write_bytes(b"x")
                (lib / "~$1Music_Library_Inventory.xlsx").write_bytes(b"x")
                self.assertEqual(cli.inventory_from(args, lib)[0], str(lib / "1Music_Library_Inventory.xlsx"))
                (lib / "Other Inventory.xlsx").write_bytes(b"x")
                self.assertIsNone(cli.inventory_from(args, lib)[0], "two candidates: do not guess")
                self.assertIsNone(cli.inventory_from(Namespace(inventory=None, no_inventory=True), lib)[0])
                self.assertEqual(cli.inventory_from(Namespace(inventory="~/x.xlsx", no_inventory=False), lib)[0],
                                 str(Path("~/x.xlsx").expanduser()))


@unittest.skipUnless(HAVE_FLAC and HAVE_FFMPEG, "flac and ffmpeg required")
class FindThenSave(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.items = [
            make_item("ta-log", "Test Artist - Test Album", "Test Artist", tracks=4, rip=True, downloads=50, seed=1),
            make_item("ta-mp3", "Test Artist - Fake Album", "Test Artist", tracks=4, rip=True, downloads=40, seed=3,
                      mp3=True, album_tag="Fake Album"),
            make_item("ta-two", "Test Artist - Second Record", "Test Artist", tracks=4, rip=True, downloads=30, seed=6,
                      album_tag="Second Record"),
        ]

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        self.out, self.lib = str(root / "runs"), root / "1Music"
        (self.lib / "Radiohead" / "OK Computer").mkdir(parents=True)
        (self.lib / "Radiohead" / "OK Computer" / "01 Airbag.flac").write_bytes(b"mine")
        self.xlsx = build_workbook(self.lib / "1Music_Library_Inventory.xlsx")
        self.original = self.xlsx.read_bytes()
        self.client = FakeIA(self.items)
        self.env = mock.patch.dict(os.environ, {}, clear=False)
        self.env.start()
        for k in ("IAFLAC_LIBRARY", "IAFLAC_INVENTORY"):
            os.environ.pop(k, None)
        self.config = root / "config.json"
        os.environ["IAFLAC_CONFIG"] = str(self.config)

    def tearDown(self):
        self.env.stop()
        self.tmp.cleanup()

    def args(self, *items, **kw):
        base = dict(items=list(items or ["Test Artist - Test Album"]), file=None, top=3, include_live=False, free_only=False,
                    deep="none", no_spectral=True, max_deep_mb=1500, out=self.out, rip_only=False, max_candidates=10,
                    search_only=False, select=None, library=str(self.lib), inventory=None, no_inventory=False, save=None,
                    yes=False, art="cover", rip_files="skip", track_names="library")
        return Namespace(**{**base, **kw})

    def album_dir(self):
        return self.lib / "Test Artist" / "Test Album"

    def rows(self):
        ws = load_workbook(self.xlsx)["Library"]
        return [tuple(r) for r in ws.iter_rows(min_row=2, max_col=5, values_only=True)]

    def prompts(self, asked):
        return [a["prompt"].strip().split("\n")[0][:30] for a in asked]

    # ---- the full flow ---------------------------------------------------------------------------

    def test_search_test_save_confirm_files_the_album_and_updates_the_inventory(self):
        code, out, err, asked = run_cli(cli.cmd_find, self.args(), self.client, answers=["1", "y"], interactive=True)
        self.assertEqual(code, 0, err)
        self.assertEqual([p.split()[0] for p in self.prompts(asked)], ["Test", "Save"], "one album: test, then one save question")
        self.assertEqual(sorted(os.listdir(self.album_dir())), ["01 Track 1.flac", "02 Track 2.flac", "03 Track 3.flac", "04 Track 4.flac"])
        self.assertFalse(any(n.endswith((".log", ".cue")) for n in os.listdir(self.album_dir())), "rip files skipped by default")
        songs = [r for r in self.rows() if r[0] == "Test Artist"]
        self.assertEqual([r[2] for r in songs], ["Track 1", "Track 2", "Track 3", "Track 4"])
        self.assertIn("(the only inventory workbook in the library folder", out)
        self.assertEqual(len(list(self.lib.glob("*.backup-*.xlsx"))), 1)

    def test_nothing_is_downloaded_before_the_save_confirmation(self):
        _, out, _, asked = run_cli(cli.cmd_find, self.args(), self.client, answers=["1", "n"], interactive=True)
        confirm = asked[-1]
        self.assertIn("Save 1 album(s)", confirm["prompt"])
        self.assertIn("What saving would do (nothing is written yet)", confirm["shown"])
        self.assertIn("destination:", confirm["shown"])
        self.assertIn("01 Track 1.flac", confirm["shown"])
        self.assertFalse(self.album_dir().exists(), "declined: nothing written")
        self.assertEqual(self.xlsx.read_bytes(), self.original)
        self.assertIn("Nothing saved.", out)
        self.assertEqual([c for c in self.client.calls if c[0] == "stream"], [], "no full file was streamed either")

    def test_the_save_prompt_comes_after_the_tests_and_the_default_is_to_save_nothing(self):
        _, out, _, asked = run_cli(cli.cmd_find, self.args(), self.client, answers=["1", ""], interactive=True)
        self.assertEqual(len(asked), 2)
        self.assertIn("y = yes", asked[1]["prompt"])
        self.assertIn("Tested albums you can save", asked[1]["shown"])
        self.assertIn("tests V0 V1", asked[1]["shown"])
        self.assertFalse(self.album_dir().exists(), "Enter is not a yes")
        self.assertIn("Nothing saved.", out)

    def test_with_several_tested_albums_you_choose_which_and_then_confirm(self):
        code, out, _, asked = run_cli(cli.cmd_find, self.args("Test Artist", select=None), self.client,
                                      answers=["a", "2", "y"], interactive=True)
        self.assertEqual(code, 0)
        self.assertIn("Save which to your library?", asked[1]["prompt"])
        self.assertEqual(len(asked), 3)
        self.assertTrue((self.lib / "Test Artist" / "Second Record" / "01 Track 1.flac").exists())
        self.assertFalse(self.album_dir().exists(), "only #2 was chosen")

    def test_only_albums_that_were_tested_can_be_saved(self):
        _, _, _, asked = run_cli(cli.cmd_find, self.args("Test Artist"), self.client, answers=["1", ""], interactive=True)
        listed = asked[1]["shown"].split("Tested albums you can save")[1]
        self.assertIn("/ta-log", listed)
        self.assertNotIn("ta-mp3", listed, "found but not tested, so not offered")

    def test_a_failed_or_suspect_tier_d_album_is_not_offered(self):
        _, out, _, asked = run_cli(cli.cmd_find, self.args("Test Artist - Fake Album", deep="all", no_spectral=False),
                                   self.client, answers=["1"], interactive=True)
        self.assertEqual(len(asked), 1, "no save prompt when nothing can be saved")
        self.assertIn("Not offered", out)
        self.assertIn("ta-mp3", out)
        self.assertIn("None of the tested albums can be saved", out)
        self.assertNotIn("Test Artist", os.listdir(self.lib))

    def test_names_can_be_edited_before_saving(self):
        code, out, _, asked = run_cli(cli.cmd_find, self.args(), self.client,
                                      answers=["1", "e", "Test Artist | Renamed Album", "y"], interactive=True)
        self.assertEqual(code, 0)
        self.assertTrue((self.lib / "Test Artist" / "Renamed Album" / "01 Track 1.flac").exists())
        self.assertFalse(self.album_dir().exists())
        self.assertIn("edited by you", out)

    def test_an_existing_artist_spelling_in_the_library_is_reused(self):
        (self.lib / "TEST ARTIST").mkdir()
        run_cli(cli.cmd_find, self.args(), self.client, answers=["1", "y"], interactive=True)
        self.assertTrue((self.lib / "TEST ARTIST" / "Test Album" / "01 Track 1.flac").exists())
        self.assertNotIn("Test Artist", os.listdir(self.lib), "no second spelling was created (exact-case listing)")

    def test_no_inventory_flag_leaves_the_workbook_alone(self):
        run_cli(cli.cmd_find, self.args(no_inventory=True), self.client, answers=["1", "y"], interactive=True)
        self.assertTrue(self.album_dir().exists())
        self.assertEqual(self.xlsx.read_bytes(), self.original)

    def test_without_a_library_it_asks_once_saves_on_yes_and_remembers_the_answer(self):
        code, out, _, asked = run_cli(cli.cmd_find, self.args(library=None), self.client,
                                      answers=["1", str(self.lib), "y"], interactive=True)
        self.assertEqual(code, 0)
        self.assertIn("Where is your music library?", asked[1]["prompt"])
        self.assertTrue((self.album_dir() / "01 Track 1.flac").exists(), "saved as soon as you said yes")
        self.assertEqual(json.loads(self.config.read_text())["library"], str(self.lib))
        # next time it does not ask again
        (self.lib / "Test Artist").rename(self.lib / "moved-away")
        _, _, _, asked = run_cli(cli.cmd_find, self.args(library=None), self.client, answers=["1", "n"], interactive=True)
        self.assertEqual([p.split()[0] for p in self.prompts(asked)], ["Test", "Save"])

    def test_a_library_that_is_not_a_folder_is_asked_again_and_s_skips_saving(self):
        _, out, _, asked = run_cli(cli.cmd_find, self.args(library=None), self.client,
                                   answers=["1", str(self.lib / "nope"), "s"], interactive=True)
        self.assertIn("is not a folder", out)
        self.assertFalse(self.album_dir().exists())
        self.assertFalse(self.config.exists(), "nothing remembered when skipped")

    def test_without_a_library_and_without_a_terminal_it_says_how_and_saves_nothing(self):
        _, out, _, asked = run_cli(cli.cmd_find, self.args(library=None, select="1"), self.client, interactive=False)
        self.assertEqual(asked, [])
        self.assertIn("--library", out)
        self.assertFalse(self.album_dir().exists())

    def test_the_library_can_come_from_the_environment(self):
        os.environ["IAFLAC_LIBRARY"] = str(self.lib)
        run_cli(cli.cmd_find, self.args(library=None), self.client, answers=["1", "y"], interactive=True)
        self.assertTrue(self.album_dir().exists())

    # ---- without a terminal ----------------------------------------------------------------------

    def test_without_a_terminal_saving_needs_both_save_and_yes(self):
        code, out, _, _ = run_cli(cli.cmd_find, self.args(select="1"), self.client, interactive=False)
        self.assertEqual(code, 0)
        self.assertIn("Use --save 1 --yes", out)
        self.assertFalse(self.album_dir().exists())
        code, out, _, _ = run_cli(cli.cmd_find, self.args(select="1", save="1"), self.client, interactive=False)
        self.assertIn("Add --yes to save", out)
        self.assertFalse(self.album_dir().exists())
        self.assertEqual(self.xlsx.read_bytes(), self.original)

    def test_save_and_yes_saves_unattended(self):
        code, out, err, asked = run_cli(cli.cmd_find, self.args(select="1", save="1", yes=True), self.client, interactive=False)
        self.assertEqual(code, 0, err)
        self.assertEqual(asked, [])
        self.assertTrue((self.album_dir() / "04 Track 4.flac").exists())
        self.assertEqual(len([r for r in self.rows() if r[0] == "Test Artist"]), 4)

    def test_search_only_never_reaches_the_save_step(self):
        _, out, _, _ = run_cli(cli.cmd_find, self.args(search_only=True, save="1", yes=True), self.client, interactive=False)
        self.assertNotIn("Save\n", out)
        self.assertFalse(self.album_dir().exists())
        self.assertEqual(file_reads(self.client), [])

    def test_saving_twice_is_harmless(self):
        run_cli(cli.cmd_find, self.args(select="1", save="1", yes=True), self.client, interactive=False)
        first = self.xlsx.read_bytes()
        code, out, _, _ = run_cli(cli.cmd_find, self.args(select="1", save="1", yes=True), self.client, interactive=False)
        self.assertEqual(code, 0)
        self.assertEqual(self.xlsx.read_bytes(), first, "already in the inventory: not written again")
        self.assertEqual(len(list(self.lib.glob("*.backup-*.xlsx"))), 1)


if __name__ == "__main__":
    unittest.main()


class ListedAlbumNames(unittest.TestCase):
    book = Namespace(library_rows=[("Madvillain", "Madvillainy 2", "a", "", ""), ("Madvillain", "Album (Deluxe)", "b", "", ""),
                                   ("Cities Aviv", "Working Title For The Album Secret Waters", "c", "", "")])

    def test_the_inventory_spelling_is_used_only_for_the_same_album(self):
        self.assertEqual(cli.listed_album(self.book, "Cities Aviv", "Working Title for the Album: Secret Waters"),
                         "Working Title For The Album Secret Waters")
        self.assertIsNone(cli.listed_album(self.book, "Madvillain", "Madvillainy"), "a number must agree both ways")
        self.assertIsNone(cli.listed_album(self.book, "Madvillain", "Album"), "an edition is a different album")
        self.assertIsNone(cli.listed_album(self.book, "Someone Else", "Madvillainy 2"))


@unittest.skipUnless(HAVE_FLAC and HAVE_FFMPEG, "flac and ffmpeg required")
class SaveAnyTestedResult(unittest.TestCase):
    """Every tested result can be saved: past the few the report lists, duplicate uploads, and albums the inventory
    already lists (a library being rebuilt)."""
    setUp, tearDown, args, rows = FindThenSave.setUp, FindThenSave.tearDown, FindThenSave.args, FindThenSave.rows

    @classmethod
    def setUpClass(cls):
        cls.items = [
            make_item("ta-log", "Test Artist - Test Album", "Test Artist", tracks=4, rip=True, downloads=50, seed=1),
            make_item("ta-dup", "Test Artist - Test Album (FLAC)", "Test Artist", tracks=4, rip=False, downloads=20, seed=2),
        ] + [make_item(f"ta-{n}", f"Test Artist - Record {n}", "Test Artist", tracks=4, rip=True, downloads=10 - i,
                       seed=10 + i, album_tag=f"Record {n}") for i, n in enumerate(["Two", "Three", "Four"])]

    def saved_list(self, asked):
        return asked[1]["shown"].split("Tested albums you can save")[1]

    def test_more_than_the_top_three_and_duplicate_uploads_are_offered(self):
        _, _, _, asked = run_cli(cli.cmd_find, self.args("Test Artist"), self.client, answers=["a", ""], interactive=True)
        listed = self.saved_list(asked)
        for ident in ("ta-log", "ta-dup", "ta-Two", "ta-Three", "ta-Four"):
            self.assertIn(f"/{ident}\n", listed + "\n")
        self.assertRegex(listed, r"Test Album\s+\(another upload of #1\)")

    def test_a_duplicate_upload_can_be_saved(self):
        _, out, _, asked = run_cli(cli.cmd_find, self.args("Test Artist - Test Album", no_inventory=True), self.client,
                                   answers=["a", "2", "y"], interactive=True)
        listed = self.saved_list(asked)
        self.assertIn("/ta-dup", listed.split("  2. ")[1], "#2 is the duplicate upload")
        self.assertTrue((self.lib / "Test Artist" / "Test Album" / "01 Track 1.flac").exists(), out)

    def test_an_album_the_inventory_already_lists_downloads_again_without_duplicate_rows(self):
        from tests.test_inventory import DATA
        at = next(i for i, r in enumerate(DATA) if r[0] == "The Marias")
        listed = [("Test Artist", "Record Two", f"Track {k}", "FLAC (16-bit/44.1kHz)", "FLAC") for k in range(1, 5)]
        build_workbook(self.xlsx, rows=DATA[:at] + listed + DATA[at:])
        before = self.rows()
        code, out, _, asked = run_cli(cli.cmd_find, self.args("Test Artist - Record Two"), self.client,
                                      answers=["1", "y"], interactive=True)
        self.assertEqual(code, 0, out)
        self.assertIn("(in your inventory: downloads again, no rows added)", asked[1]["shown"])
        self.assertIn("ALREADY in the inventory: the files are downloaded", asked[1]["shown"])
        self.assertEqual(sorted(os.listdir(self.lib / "Test Artist" / "Record Two"))[:1], ["01 Track 1.flac"])
        self.assertEqual(self.rows(), before, "no duplicate rows")
